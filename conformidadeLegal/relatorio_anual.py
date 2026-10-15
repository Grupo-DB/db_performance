"""
Relatório Anual de Acompanhamento das Atividades — o documento que a FEPAM pede
todo ano para cada licença: as condicionantes na ordem da licença, cada uma com a
resposta da empresa e as fotos que comprovam.

Montado com o que o módulo já guarda:
  - a licença é uma `Norma` (tipo LICENCA/OUTORGA/TAC);
  - cada condicionante é um `Requisito` (referência = número da condicionante,
    trecho_fonte = texto literal, secao = "Quanto à Atividade");
  - a resposta é a `Verificacao` mais recente DENTRO do período do relatório, e as
    fotos são os anexos de imagem dela (legenda = "Figura N: …").

Sai em Word (para a equipe ajustar antes de protocolar) e em PDF, com o mesmo
conteúdo. Capa e apresentação vêm de `Norma.relatorio_anual`, com padrões tirados
da própria licença.
"""
import io
import logging
import os
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from xml.sax.saxutils import escape

from dateutil.relativedelta import relativedelta
from django.db.models import Prefetch

from .models import EvidenciaAnexo, Requisito, Verificacao

logger = logging.getLogger(__name__)

MESES = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro',
         'novembro', 'dezembro']
EXTENSOES_IMAGEM = ('.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp', '.tif', '.tiff')
# Foto de celular tem 4000 px e 5 MB: no documento ela ocupa 15 cm. Reduzir aqui é o que
# impede um relatório de 40 fotos de virar um arquivo de 200 MB que o SOL recusa.
LADO_MAXIMO_FOTO = 1600

CAMPOS_CONFIG = ('titulo', 'empresa', 'cnpj', 'municipio', 'atividade', 'processos', 'licenca', 'item',
                 'apresentacao', 'cabecalho', 'data_capa', 'logo')
# Listas de ids de requisito guardadas junto da configuração.
LISTAS_CONFIG = ('sem_resposta', 'fora')
LOGO_PADRAO = os.path.join(os.path.dirname(__file__), 'recursos', 'logo_relatorio.png')

APRESENTACAO_PADRAO = (
    '{empresa}, empresa inscrita no CNPJ sob nº {cnpj}, vem apresentar o Relatório Técnico Anual comprovando o '
    'atendimento do item {item} e seus subitens, das condições e restrições da {licenca}. O presente Relatório '
    'compreende as atividades desenvolvidas entre {inicio} e {fim}.'
)


# ── Configuração (capa, apresentação, cabeçalho) ─────────────────────────────

def configuracao_padrao(norma) -> dict:
    numero = norma.numero + (f'/{norma.ano}' if norma.ano and norma.numero and '/' not in norma.numero else '')
    licenca = f'Licença de Operação nº {numero}' if numero else norma.identificacao
    return {
        'titulo': 'Relatório Anual de Acompanhamento das Atividades',
        'empresa': 'Dagoberto Barcellos S/A',
        'cnpj': '87.678.934/0001-65',
        'municipio': _municipio(norma) or 'Caçapava do Sul – RS',
        'atividade': '',
        'processos': '',
        'licenca': licenca,
        'item': '',
        'apresentacao': APRESENTACAO_PADRAO,
        'cabecalho': '{titulo} – ' + (f'{norma.orgao} – ' if norma.orgao else '') + '{licenca}',
        'data_capa': '',
        'logo': 'sim',
    }


def _municipio(norma) -> str:
    """Cidade da unidade da licença (o cadastro de filiais tem cidade e UF)."""
    u = getattr(norma, 'unidade', None)
    if u and u.cidade:
        return f'{u.cidade} – {u.estado}' if u.estado else u.cidade
    return ''


def referencia_curta(ref: str) -> str:
    """'Item 1.4' / 'Condicionante 2.2' → '1.4': no relatório o número fala sozinho."""
    return re.sub(r'^\s*(?:item|condicionante|cond\.|condi[çc][ãa]o)\s+', '', ref or '', flags=re.IGNORECASE).strip()


def configuracao(norma) -> dict:
    """Padrões da licença com o que a equipe já salvou por cima (campo vazio volta ao padrão)."""
    cfg = configuracao_padrao(norma)
    salvo = norma.relatorio_anual or {}
    for campo in CAMPOS_CONFIG:
        valor = salvo.get(campo)
        if isinstance(valor, str) and valor.strip():
            cfg[campo] = valor.strip()
    return cfg


def limpar_configuracao(dados: dict) -> dict:
    """Só os campos conhecidos, como texto (e as listas de ids: sem resposta própria, fora do relatório)."""
    limpo = {c: str(dados.get(c) or '').strip()[:4000] for c in CAMPOS_CONFIG if c in dados}
    for lista in LISTAS_CONFIG:
        if lista in dados:
            try:
                limpo[lista] = sorted({int(i) for i in (dados.get(lista) or [])})
            except (TypeError, ValueError):
                limpo[lista] = []
    return limpo


def lista_ids(norma, nome: str) -> list[int]:
    valor = (norma.relatorio_anual or {}).get(nome) or []
    return [i for i in valor if isinstance(i, int)]


def sem_resposta_propria(norma) -> list[int]:
    return lista_ids(norma, 'sem_resposta')


def caminho_logo(norma, cfg: dict) -> str | None:
    """Logo próprio da licença, senão o do grupo; None quando a capa vai sem logo."""
    if cfg.get('logo') == 'nao':
        return None
    if norma.relatorio_logo:
        try:
            if os.path.exists(norma.relatorio_logo.path):
                return norma.relatorio_logo.path
        except (ValueError, NotImplementedError):
            pass
    return LOGO_PADRAO if os.path.exists(LOGO_PADRAO) else None


class _Molde(dict):
    """format_map que deixa {desconhecido} como está, em vez de explodir com KeyError."""

    def __missing__(self, chave):
        return '{' + chave + '}'


def preencher(texto: str, valores: dict) -> str:
    try:
        return texto.format_map(_Molde(valores))
    except (ValueError, IndexError):
        # Chave solta ("{" sem fechar) no texto digitado: imprime como foi escrito.
        return texto


def mes_ano(d: date) -> str:
    return f'{MESES[d.month - 1]} de {d.year}'


def periodo_padrao(hoje: date | None = None) -> tuple[date, date]:
    """Os 12 meses fechados antes do mês atual (o relatório de dezembro cobre dez → nov)."""
    hoje = hoje or date.today()
    fim = hoje.replace(day=1) - timedelta(days=1)
    inicio = (fim + timedelta(days=1)) - relativedelta(years=1)
    return inicio, fim


# ── Ordem das condicionantes ─────────────────────────────────────────────────

def chave_referencia(ref: str) -> tuple:
    """
    "5.13.2" → (0, 5, 13, 2). Ordem numérica, e não alfabética: na alfabética o
    5.10 vem antes do 5.2. Referência sem número vai para o fim.
    """
    numeros = [int(n) for n in re.findall(r'\d+', ref or '')]
    return (0, *numeros) if numeros else (1,)


def _titulo_secao(secao: str, primeira_ref: str) -> str:
    """Põe o número do capítulo ("5. Quanto à Atividade") quando o título veio sem ele."""
    secao = secao.strip().rstrip(':')
    if not secao or re.match(r'^\d', secao):
        return secao
    numeros = re.findall(r'\d+', primeira_ref or '')
    return f'{numeros[0]}. {secao}' if numeros else secao


# ── Montagem (dados comuns à prévia, ao Word e ao PDF) ───────────────────────

@dataclass
class Foto:
    id: int
    nome: str
    legenda: str
    caminho: str | None
    url: str
    numero: int = 0


@dataclass
class Item:
    requisito_id: int
    referencia: str
    texto: str
    aplicabilidade: str
    resposta: str = ''
    origem_resposta: str = ''  # VERIFICACAO | JUSTIFICATIVA | '' (pendente)
    resultado: str = ''
    data_resposta: date | None = None
    verificacao_id: int | None = None
    fotos: list = field(default_factory=list)
    documentos: list = field(default_factory=list)
    anterior: dict | None = None
    # Subitem respondido junto no item seguinte (3.1.1…3.1.4 → resposta no 3.1.5): sai
    # só com o texto, sem o aviso de "preencher".
    sem_resposta_propria: bool = False
    # Desmarcado na tela: a condicionante não sai no documento (ex.: "Disposições finais").
    no_relatorio: bool = True

    @property
    def pendente(self) -> bool:
        return self.no_relatorio and not self.resposta.strip() and not self.sem_resposta_propria


@dataclass
class Secao:
    titulo: str
    chave: str = ''
    itens: list = field(default_factory=list)


def _e_imagem(nome: str) -> bool:
    return nome.lower().endswith(EXTENSOES_IMAGEM)


def _caminho(anexo) -> str | None:
    try:
        return anexo.arquivo.path
    except (ValueError, NotImplementedError):
        return None


def montar(norma, inicio: date, fim: date, url_de=None) -> dict:
    """
    Tudo o que o relatório precisa, já ordenado e com as figuras numeradas.
    `url_de(anexo)` monta a URL absoluta das fotos para a prévia na tela.
    """
    url_de = url_de or (lambda a: a.arquivo.url)
    verificacoes = Verificacao.objects.order_by('-data', '-id').prefetch_related(
        Prefetch('anexos', queryset=EvidenciaAnexo.objects.order_by('id')))
    requisitos = sorted(
        Requisito.objects.filter(norma=norma, ativo=True).prefetch_related(Prefetch('verificacoes', queryset=verificacoes)),
        key=lambda r: (chave_referencia(r.referencia), r.id),
    )

    sem_resposta = set(sem_resposta_propria(norma))
    fora = set(lista_ids(norma, 'fora'))
    secoes: list[Secao] = []
    atual: Secao | None = None
    figura = 0
    for req in requisitos:
        item = Item(
            requisito_id=req.id, referencia=referencia_curta(req.referencia),
            texto=(req.trecho_fonte or '').strip() or req.descricao.strip(),
            aplicabilidade=req.aplicabilidade, sem_resposta_propria=req.id in sem_resposta,
            no_relatorio=req.id not in fora,
        )
        no_periodo = next((v for v in req.verificacoes.all() if inicio <= v.data <= fim), None)
        if no_periodo:
            item.resposta = no_periodo.evidencia.strip()
            item.origem_resposta = 'VERIFICACAO'
            item.resultado = no_periodo.resultado
            item.data_resposta = no_periodo.data
            item.verificacao_id = no_periodo.id
            for a in no_periodo.anexos.all():
                destino = item.fotos if _e_imagem(a.nome or a.arquivo.name) else item.documentos
                destino.append(Foto(id=a.id, nome=a.nome or os.path.basename(a.arquivo.name), legenda=a.legenda,
                                    caminho=_caminho(a), url=url_de(a)))
        elif req.aplicabilidade == 'NAO_APLICAVEL' and req.justificativa.strip():
            item.resposta = req.justificativa.strip()
            item.origem_resposta = 'JUSTIFICATIVA'
        if item.pendente:
            # Sem resposta no período: a última de antes vira sugestão de "repetir".
            antes = next((v for v in req.verificacoes.all() if v.data < inicio and v.evidencia.strip()), None)
            if antes:
                item.anterior = {
                    'verificacao': antes.id, 'data': antes.data.isoformat(), 'resultado': antes.resultado,
                    'evidencia': antes.evidencia.strip(), 'fotos': sum(1 for a in antes.anexos.all()
                                                                       if _e_imagem(a.nome or a.arquivo.name)),
                }
        if item.no_relatorio:
            for foto in item.fotos:
                figura += 1
                foto.numero = figura

        secao = (req.secao or '').strip()
        if atual is None or (secao and secao.lower() != atual.chave):
            atual = Secao(titulo=_titulo_secao(secao, req.referencia), chave=secao.lower())
            secoes.append(atual)
        atual.itens.append(item)

    todos = [i for s in secoes for i in s.itens]
    itens = [i for i in todos if i.no_relatorio]
    cfg = configuracao(norma)
    valores = {**cfg, 'inicio': mes_ano(inicio), 'fim': mes_ano(fim)}
    return {
        'norma': norma,
        'inicio': inicio,
        'fim': fim,
        'config': cfg,
        'apresentacao': preencher(cfg['apresentacao'], valores),
        'cabecalho': preencher(cfg['cabecalho'], valores),
        'data_capa': cfg['data_capa'] or f'{MESES[date.today().month - 1]} {date.today().year}'.upper(),
        'secoes': secoes,
        'logo': caminho_logo(norma, cfg),
        'totais': {
            'itens': len(itens),
            'fora': len(todos) - len(itens),
            'respondidos': sum(1 for i in itens if i.resposta.strip()),
            'pendentes': sum(1 for i in itens if i.pendente),
            'com_anterior': sum(1 for i in itens if i.pendente and i.anterior),
            'fotos': figura,
            'sem_legenda': sum(1 for i in itens for f in i.fotos if not f.legenda.strip()),
        },
    }


def para_json(dados: dict) -> dict:
    """A prévia da tela: o mesmo conteúdo do documento, mais o que falta preencher."""
    def foto(f: Foto):
        return {'id': f.id, 'nome': f.nome, 'legenda': f.legenda, 'url': f.url, 'numero': f.numero}

    return {
        'inicio': dados['inicio'].isoformat(),
        'fim': dados['fim'].isoformat(),
        'config': dados['config'],
        'padrao': configuracao_padrao(dados['norma']),
        'logo_proprio': dados['norma'].relatorio_logo.url if dados['norma'].relatorio_logo else None,
        'apresentacao': dados['apresentacao'],
        'cabecalho': dados['cabecalho'],
        'data_capa': dados['data_capa'],
        'totais': dados['totais'],
        'secoes': [{
            'titulo': s.titulo,
            'itens': [{
                'requisito': i.requisito_id, 'referencia': i.referencia, 'texto': i.texto,
                'aplicabilidade': i.aplicabilidade, 'resposta': i.resposta, 'origem_resposta': i.origem_resposta,
                'resultado': i.resultado, 'data_resposta': i.data_resposta.isoformat() if i.data_resposta else None,
                'verificacao': i.verificacao_id, 'pendente': i.pendente, 'anterior': i.anterior,
                'sem_resposta_propria': i.sem_resposta_propria, 'no_relatorio': i.no_relatorio,
                'fotos': [foto(f) for f in i.fotos], 'documentos': [foto(f) for f in i.documentos],
            } for i in s.itens],
        } for s in dados['secoes']],
    }


def secoes_do_documento(dados: dict) -> list:
    """Capítulos com só as condicionantes marcadas para o relatório (capítulo vazio some)."""
    saida = []
    for s in dados['secoes']:
        itens = [i for i in s.itens if i.no_relatorio]
        if itens:
            saida.append(Secao(titulo=s.titulo, chave=s.chave, itens=itens))
    return saida


def nome_arquivo(dados: dict, extensao: str) -> str:
    licenca = re.sub(r'[^0-9A-Za-z]+', '_', dados['config']['licenca']).strip('_')[:60] or 'licenca'
    return f'Relatorio_Anual_{licenca}_{dados["fim"]:%Y-%m}.{extensao}'


def _legenda(f: Foto) -> str:
    texto = f.legenda.strip().rstrip('.')
    return f'Figura {f.numero}: {texto}.' if texto else f'Figura {f.numero}.'


def _foto_jpeg(f: Foto) -> tuple[io.BytesIO, int, int] | None:
    """A foto em pé (EXIF), reduzida e em JPEG; None se o arquivo sumiu ou não abre."""
    if not f.caminho or not os.path.exists(f.caminho):
        return None
    try:
        from PIL import Image, ImageOps
        with Image.open(f.caminho) as img:
            img = ImageOps.exif_transpose(img)
            if img.mode not in ('RGB', 'L'):
                img = img.convert('RGB')
            img.thumbnail((LADO_MAXIMO_FOTO, LADO_MAXIMO_FOTO))
            saida = io.BytesIO()
            img.save(saida, format='JPEG', quality=85, optimize=True)
            saida.seek(0)
            return saida, img.width, img.height
    except Exception:  # noqa: BLE001 — uma foto corrompida não pode derrubar o relatório
        logger.warning('Relatório anual: não foi possível ler a foto %s', f.caminho, exc_info=True)
        return None


def _logo(caminho: str | None) -> tuple[io.BytesIO, int, int] | None:
    """Logo em PNG (mantém o fundo transparente), reduzido; None se não houver ou não abrir."""
    if not caminho or not os.path.exists(caminho):
        return None
    try:
        from PIL import Image
        with Image.open(caminho) as img:
            img.thumbnail((1200, 1200))
            if img.mode not in ('RGB', 'RGBA', 'L', 'LA'):
                img = img.convert('RGBA')
            saida = io.BytesIO()
            img.save(saida, format='PNG')
            saida.seek(0)
            return saida, img.width, img.height
    except Exception:  # noqa: BLE001
        logger.warning('Relatório anual: logo %s não abriu', caminho, exc_info=True)
        return None


def _medidas(largura: int, altura: int, max_l: float, max_a: float) -> tuple[float, float]:
    escala = min(max_l / largura, max_a / altura)
    return largura * escala, altura * escala


TEXTO_PENDENTE = '[Sem resposta registrada no período — preencher]'


# ── Word ─────────────────────────────────────────────────────────────────────

def gerar_docx(dados: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    cfg = dados['config']
    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = 'Arial'
    estilo.font.size = Pt(11)
    estilo.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.line_spacing = 1.25

    secao = doc.sections[0]
    secao.page_height, secao.page_width = Cm(29.7), Cm(21)
    secao.left_margin = secao.top_margin = Cm(3)
    secao.right_margin = secao.bottom_margin = Cm(2)
    secao.different_first_page_header_footer = True
    largura_util = 21 - 3 - 2

    # Cabeçalho: número da página à direita. Rodapé: a linha da licença, como no modelo.
    cab = secao.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _campo_pagina(cab, OxmlElement, qn)
    rod = secao.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rod.add_run(dados['cabecalho'])
    r.font.size = Pt(9)
    r.italic = True
    _borda_superior(rod, OxmlElement, qn)

    def par(texto='', negrito=False, tamanho=None, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=None,
            primeira=None, antes=None, depois=None, manter=False):
        p = doc.add_paragraph()
        p.alignment = alinhamento
        f = p.paragraph_format
        if recuo is not None:
            f.left_indent = Cm(recuo)
        if primeira is not None:
            f.first_line_indent = Cm(primeira)
        if antes is not None:
            f.space_before = Pt(antes)
        if depois is not None:
            f.space_after = Pt(depois)
        if manter:
            f.keep_with_next = True
        if texto:
            run = p.add_run(texto)
            run.bold = negrito
            if tamanho:
                run.font.size = Pt(tamanho)
        return p

    # Capa
    centro = WD_ALIGN_PARAGRAPH.CENTER
    logo = _logo(dados.get('logo'))
    if logo:
        conteudo, larg, alt = logo
        w, hh = _medidas(larg, alt, 6, 4)
        par(alinhamento=centro, antes=24, depois=30).add_run().add_picture(conteudo, width=Cm(w), height=Cm(hh))
    else:
        par(antes=90)
    par(cfg['titulo'].upper(), negrito=True, tamanho=16, alinhamento=centro, depois=60)
    par(cfg['empresa'].upper(), negrito=True, tamanho=14, alinhamento=centro)
    par(cfg['municipio'].upper(), negrito=True, tamanho=12, alinhamento=centro, depois=60)
    if cfg['atividade']:
        par(cfg['atividade'].upper(), negrito=True, tamanho=12, alinhamento=centro)
    for linha in cfg['processos'].splitlines():
        if linha.strip():
            par(linha.strip().upper(), negrito=True, tamanho=12, alinhamento=centro, depois=0)
    par(cfg['licenca'].upper(), negrito=True, tamanho=12, alinhamento=centro, antes=6)
    par(antes=140)
    par(dados['data_capa'].upper(), negrito=True, tamanho=12, alinhamento=centro)
    doc.paragraphs[-1].runs[0].add_break(WD_BREAK.PAGE)

    # Apresentação
    par('APRESENTAÇÃO', negrito=True, antes=6, depois=12, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    for linha in dados['apresentacao'].splitlines():
        if linha.strip():
            par(linha.strip(), primeira=2)
    titulo = f'ATENDIMENTO DAS CONDIÇÕES E RESTRIÇÕES DA {cfg["licenca"].upper()}'
    par(titulo, negrito=True, antes=18, depois=12, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, manter=True)

    for s in secoes_do_documento(dados):
        if s.titulo:
            par(s.titulo + ':', negrito=True, antes=12, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, manter=True)
        for item in s.itens:
            # Condicionante em negrito com recuo deslocado: "5.13.2-  texto…"
            p = par(recuo=1.25, primeira=-1.25, antes=6, manter=True)
            if item.referencia:
                rr = p.add_run(f'{item.referencia}-\t')
                rr.bold = True
            p.add_run(item.texto).bold = True
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.25))
            if item.pendente:
                aviso = par(recuo=1.25)
                run = aviso.add_run(TEXTO_PENDENTE)
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            else:
                for linha in item.resposta.splitlines():
                    if linha.strip():
                        par(linha.strip(), recuo=1.25)
            for d in item.documentos:
                p = par(recuo=1.25)
                run = p.add_run(f'Anexo: {d.nome}')
                run.italic = True
                run.font.size = Pt(10)
            for f in item.fotos:
                imagem = _foto_jpeg(f)
                if imagem is None:
                    p = par(f'[Foto não encontrada: {f.nome}]', alinhamento=centro)
                    p.runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                else:
                    conteudo, larg, alt = imagem
                    w, h = _medidas(larg, alt, largura_util - 1, 11)
                    p = par(alinhamento=centro, antes=6, depois=2, manter=True)
                    p.add_run().add_picture(conteudo, width=Cm(w), height=Cm(h))
                leg = par(_legenda(f), alinhamento=centro, depois=10)
                leg.runs[0].font.size = Pt(10)
                leg.runs[0].italic = True

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _campo_pagina(paragrafo, OxmlElement, qn):
    """Campo PAGE do Word (o número é calculado ao abrir, como num documento feito à mão)."""
    run = paragrafo.add_run()
    for tipo, texto in (('begin', None), (None, 'PAGE'), ('separate', None), (None, '1'), ('end', None)):
        if tipo:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tipo)
        elif texto == 'PAGE':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = ' PAGE '
        else:
            el = OxmlElement('w:t')
            el.text = texto
        run._r.append(el)


def _borda_superior(paragrafo, OxmlElement, qn):
    ppr = paragrafo._p.get_or_add_pPr()
    bordas = OxmlElement('w:pBdr')
    topo = OxmlElement('w:top')
    for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '4'), ('w:color', '808080')):
        topo.set(qn(k), v)
    bordas.append(topo)
    ppr.append(bordas)


# ── PDF ──────────────────────────────────────────────────────────────────────

_FONTES = None


def _fontes() -> tuple[str, str, str]:
    """
    (normal, negrito, itálico). Fonte TrueType quando o servidor tem uma — a Helvetica
    embutida do reportlab só desenha o que cabe no WinAnsi, e um "≥" digitado na
    resposta viraria um quadrado preto.
    """
    global _FONTES
    if _FONTES:
        return _FONTES
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidatos = [
        ('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
         '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
         '/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf'),
        ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
         '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
         '/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf'),
    ]
    for normal, negrito, italico in candidatos:
        if all(os.path.exists(c) for c in (normal, negrito, italico)):
            try:
                pdfmetrics.registerFont(TTFont('RA-Normal', normal))
                pdfmetrics.registerFont(TTFont('RA-Negrito', negrito))
                pdfmetrics.registerFont(TTFont('RA-Italico', italico))
                from reportlab.lib.fonts import addMapping
                addMapping('RA-Normal', 0, 0, 'RA-Normal')
                addMapping('RA-Normal', 1, 0, 'RA-Negrito')
                addMapping('RA-Normal', 0, 1, 'RA-Italico')
                addMapping('RA-Normal', 1, 1, 'RA-Negrito')
                _FONTES = ('RA-Normal', 'RA-Negrito', 'RA-Italico')
                return _FONTES
            except Exception:  # noqa: BLE001
                logger.warning('Relatório anual: fonte %s não registrou', normal, exc_info=True)
    _FONTES = ('Helvetica', 'Helvetica-Bold', 'Helvetica-Oblique')
    return _FONTES


def _xml(texto: str) -> str:
    texto = escape(texto)
    if _fontes()[0] == 'Helvetica':
        texto = texto.encode('cp1252', errors='replace').decode('cp1252')
    return texto


def gerar_pdf(dados: dict) -> bytes:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer)

    normal, negrito, italico = _fontes()
    cfg = dados['config']
    base = ParagraphStyle('base', fontName=normal, fontSize=11, leading=15, alignment=TA_JUSTIFY, spaceAfter=5)
    capa = ParagraphStyle('capa', parent=base, fontName=negrito, alignment=TA_CENTER, fontSize=12, leading=17)
    titulo = ParagraphStyle('titulo', parent=base, fontName=negrito, alignment=TA_LEFT, spaceBefore=10, spaceAfter=8)
    cond = ParagraphStyle('cond', parent=base, fontName=negrito, leftIndent=1.25 * cm, bulletIndent=0,
                          bulletFontName=negrito, bulletFontSize=11, spaceBefore=6)
    resp = ParagraphStyle('resp', parent=base, leftIndent=1.25 * cm)
    legenda = ParagraphStyle('legenda', parent=base, fontName=italico, fontSize=10, leading=13, alignment=TA_CENTER,
                             spaceAfter=10)
    anexo = ParagraphStyle('anexo', parent=resp, fontName=italico, fontSize=10)
    pendente = ParagraphStyle('pendente', parent=resp, textColor='#C00000', backColor='#FFF2A8')

    def rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont(normal, 10)
        canvas.drawRightString(A4[0] - 2 * cm, A4[1] - 1.6 * cm, str(doc.page))
        canvas.setStrokeColor('#808080')
        canvas.setLineWidth(0.5)
        canvas.line(3 * cm, 1.75 * cm, A4[0] - 2 * cm, 1.75 * cm)
        canvas.setFont(italico, 8.5)
        canvas.drawCentredString((A4[0] + 1 * cm) / 2, 1.3 * cm, dados['cabecalho'])
        canvas.restoreState()

    saida = io.BytesIO()
    doc = SimpleDocTemplate(saida, pagesize=A4, leftMargin=3 * cm, rightMargin=2 * cm, topMargin=2.5 * cm,
                            bottomMargin=2.3 * cm, title=f'{cfg["titulo"]} – {cfg["licenca"]}', author=cfg['empresa'])
    largura_util = A4[0] - 5 * cm
    h = []

    logo = _logo(dados.get('logo'))
    if logo:
        conteudo, larg, alt = logo
        w, hh = _medidas(larg, alt, 6 * cm, 4 * cm)
        h += [Spacer(1, 0.5 * cm), Image(conteudo, width=w, height=hh, mask='auto'), Spacer(1, 1.2 * cm)]
    else:
        h.append(Spacer(1, 4 * cm))
    h += [Paragraph(_xml(cfg['titulo'].upper()), ParagraphStyle('capa1', parent=capa, fontSize=16, leading=22)),
          Spacer(1, 2 * cm),
          Paragraph(_xml(cfg['empresa'].upper()), ParagraphStyle('capa2', parent=capa, fontSize=14, leading=19)),
          Paragraph(_xml(cfg['municipio'].upper()), capa), Spacer(1, 2 * cm)]
    if cfg['atividade']:
        h.append(Paragraph(_xml(cfg['atividade'].upper()), capa))
    h += [Paragraph(_xml(l.strip().upper()), capa) for l in cfg['processos'].splitlines() if l.strip()]
    h += [Paragraph(_xml(cfg['licenca'].upper()), capa), Spacer(1, 5 * cm),
          Paragraph(_xml(dados['data_capa'].upper()), capa), PageBreak()]

    h.append(Paragraph('APRESENTAÇÃO', titulo))
    h += [Paragraph(_xml(l.strip()), ParagraphStyle('apr', parent=base, firstLineIndent=2 * cm))
          for l in dados['apresentacao'].splitlines() if l.strip()]
    h.append(Paragraph(_xml(f'ATENDIMENTO DAS CONDIÇÕES E RESTRIÇÕES DA {cfg["licenca"].upper()}'),
                       ParagraphStyle('t2', parent=titulo, spaceBefore=16)))

    for s in secoes_do_documento(dados):
        bloco_secao = [Paragraph(_xml(s.titulo + ':'), titulo)] if s.titulo else []
        for n, item in enumerate(s.itens):
            ref = f'{_xml(item.referencia)}-' if item.referencia else None
            corpo = [Paragraph(_xml(item.texto).replace('\n', '<br/>'), cond, bulletText=ref)]
            if item.pendente:
                corpo.append(Paragraph(_xml(TEXTO_PENDENTE), pendente))
            else:
                corpo += [Paragraph(_xml(l.strip()), resp) for l in item.resposta.splitlines() if l.strip()]
            corpo += [Paragraph(_xml(f'Anexo: {d.nome}'), anexo) for d in item.documentos]
            # Título do capítulo e a primeira condicionante andam juntos (não fica título órfão no pé da página).
            if n == 0 and bloco_secao:
                corpo = bloco_secao + corpo
            h.append(KeepTogether(corpo[:3]))
            h += corpo[3:]
            for f in item.fotos:
                imagem = _foto_jpeg(f)
                if imagem is None:
                    figura = Paragraph(_xml(f'[Foto não encontrada: {f.nome}]'), pendente)
                else:
                    conteudo, larg, alt = imagem
                    w, hh = _medidas(larg, alt, largura_util - 1 * cm, 11 * cm)
                    figura = Image(conteudo, width=w, height=hh)
                h.append(KeepTogether([Spacer(1, 4), figura, Spacer(1, 3), Paragraph(_xml(_legenda(f)), legenda)]))

    doc.build(h, onFirstPage=lambda c, d: None, onLaterPages=rodape)
    return saida.getvalue()
