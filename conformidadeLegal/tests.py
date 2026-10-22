import io
import shutil
import tempfile
from datetime import date, timedelta

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
from rest_framework.test import APITestCase

from kanban.models import KanbanBoard, KanbanColumn, KanbanTask

from .models import EvidenciaAnexo, Norma, PalavraChaveMonitor, PlanoAcao, Requisito, Verificacao

# Anexos dos testes vão para uma pasta descartável, não para o media/ do projeto.
MEDIA_TESTE = tempfile.mkdtemp(prefix='conformidade-testes-')


def _filiais():
    """Matriz e a filial de Almirante Tamandaré, como no cadastro de produção."""
    from avaliacoes.management.models import Empresa, Filial
    e = Empresa.objects.create(nome='Dagoberto Barcellos', cnpj='87.678.934/0001-65', endereco='BR 392',
                               cidade='Caçapava do Sul', estado='RS', codigo='01')
    comum = dict(empresa=e, endereco='x')
    return (Filial.objects.create(nome='Matriz', cnpj='87.678.934/0001-65', cidade='Caçapava do Sul', estado='RS',
                                  codigo='01', **comum),
            Filial.objects.create(nome='F08 - UP ATM', cnpj='87.678.934/0008-31', cidade='Almirante Tamandaré',
                                  estado='PR', codigo='08', **comum))


def tearDownModule():
    shutil.rmtree(MEDIA_TESTE, ignore_errors=True)


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class ConformidadeLegalTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        r = self.client.post('/conformidadeLegal/normas/', {
            'tipo': 'RESOLUCAO', 'orgao': 'CONAMA', 'numero': '430', 'ano': 2011,
            'esfera': 'FEDERAL', 'tema': 'AMBIENTAL', 'ementa': 'Condições e padrões de lançamento de efluentes.',
        }, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['identificacao'], 'Resolução CONAMA nº 430/2011')
        self.norma = Norma.objects.get(pk=r.data['id'])

    def _requisito(self, **extra):
        r = self.client.post('/conformidadeLegal/requisitos/', {
            'norma': self.norma.id, 'referencia': 'Art. 16', 'descricao': 'Monitorar o efluente.',
            'tema': 'AMBIENTAL', 'aplicabilidade': 'APLICAVEL', 'periodicidade_meses': 6, **extra,
        }, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        return Requisito.objects.get(pk=r.data['id'])

    def test_edicao_em_lote_muda_so_os_campos_enviados(self):
        matriz, atm = _filiais()
        a = self._requisito(aplicabilidade='EM_ANALISE', unidade=matriz.id)
        b = self._requisito(aplicabilidade='EM_ANALISE', unidade=atm.id)
        r = self.client.post('/conformidadeLegal/requisitos/editar-lote/', {
            'ids': [a.id, b.id],
            'campos': {'aplicabilidade': 'NAO_APLICAVEL', 'justificativa': 'Sem caldeira.', 'norma': 999},
        }, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        for req, unidade in ((a, matriz.id), (b, atm.id)):
            req.refresh_from_db()
            # `norma` não é campo de lote e é ignorado; a unidade de cada um fica.
            self.assertEqual((req.aplicabilidade, req.justificativa, req.unidade_id, req.norma_id),
                             ('NAO_APLICAVEL', 'Sem caldeira.', unidade, self.norma.id))

    def test_edicao_em_lote_e_tudo_ou_nada(self):
        a = self._requisito()
        r = self.client.post('/conformidadeLegal/requisitos/editar-lote/',
                             {'ids': [a.id, 987654], 'campos': {'justificativa': 'X'}}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.post('/conformidadeLegal/requisitos/editar-lote/',
                             {'ids': [a.id], 'campos': {'aplicabilidade': 'TALVEZ'}}, format='json')
        self.assertEqual(r.status_code, 400)
        a.refresh_from_db()
        self.assertEqual((a.justificativa, a.aplicabilidade), ('', 'APLICAVEL'))

    def test_verificacao_em_lote_compartilha_o_anexo(self):
        import os
        a, b, c = self._requisito(), self._requisito(), self._requisito()
        laudo = SimpleUploadedFile('laudo_chamine.pdf', b'%PDF-1.4 laudo', content_type='application/pdf')
        r = self.client.post('/conformidadeLegal/requisitos/verificar-lote/', {
            'ids': [a.id, b.id, c.id], 'resultado': 'ATENDIDO', 'data': '2026-10-01',
            'evidencia': 'Laudo de amostragem 2026.', 'arquivos': [laudo],
        }, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data['verificacoes'], r.data['planos']), (3, 0))
        anexos = list(EvidenciaAnexo.objects.order_by('id'))
        self.assertEqual(len(anexos), 3)
        self.assertEqual(len({x.arquivo.name for x in anexos}), 1)  # gravado uma vez só
        caminho = anexos[0].arquivo.path
        for req in (a, b, c):
            req.refresh_from_db()
            self.assertEqual((req.situacao, req.proxima_verificacao), ('ATENDIDO', date(2027, 4, 1)))

        # Apagar uma das verificações não pode levar o arquivo das outras.
        self.client.delete(f'/conformidadeLegal/verificacoes/{anexos[0].verificacao_id}/')
        self.assertTrue(os.path.exists(caminho))
        for x in anexos[1:]:
            self.client.delete(f'/conformidadeLegal/verificacoes/{x.verificacao_id}/')
        self.assertFalse(os.path.exists(caminho))

    def test_verificacao_em_lote_regras(self):
        a, b = self._requisito(), self._requisito()
        em_analise = self._requisito(aplicabilidade='EM_ANALISE')
        url = '/conformidadeLegal/requisitos/verificar-lote/'
        r = self.client.post(url, {'ids': [a.id, em_analise.id], 'resultado': 'ATENDIDO'}, format='multipart')
        self.assertEqual(r.status_code, 400)
        self.assertIn('Aplicável', str(r.data['ids']))
        # Não atendido sem plano é recusado; com plano, cada um ganha o seu.
        r = self.client.post(url, {'ids': [a.id, b.id], 'resultado': 'NAO_ATENDIDO'}, format='multipart')
        self.assertEqual(r.status_code, 400)
        self.assertIn('plano_descricao', r.data)
        r = self.client.post(url, {'ids': [a.id, b.id], 'resultado': 'NAO_ATENDIDO',
                                   'plano_descricao': 'Contratar amostragem.'}, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['planos'], 2)
        # De novo: os dois já têm plano aberto, então nada de plano duplicado.
        r = self.client.post(url, {'ids': [a.id, b.id], 'resultado': 'PARCIAL',
                                   'plano_descricao': 'Outro.'}, format='multipart')
        self.assertEqual((r.status_code, r.data['planos']), (201, 0))
        self.assertEqual(PlanoAcao.objects.count(), 2)
        self.assertFalse(Verificacao.objects.filter(requisito=em_analise).exists())

    def test_responsaveis_traz_ativos_com_o_nome_mais_humano_que_houver(self):
        User.objects.create_user('jsilva', first_name='João', last_name='Silva', password='x')
        User.objects.create_user('inativo', password='x', is_active=False)
        r = self.client.get('/conformidadeLegal/responsaveis/')
        self.assertEqual(r.status_code, 200)
        nomes = {u['login']: u['nome'] for u in r.data}
        self.assertEqual(nomes['jsilva'], 'João Silva')  # sem colaborador ligado: nome do usuário
        self.assertEqual(nomes['sgi'], 'sgi')            # sem nome nenhum: o login
        self.assertNotIn('inativo', nomes)

    def test_nome_do_colaborador_vence_o_do_usuario(self):
        from types import SimpleNamespace
        from .serializers import nome_usuario
        usuario = SimpleNamespace(username='jsilva', get_full_name=lambda: 'J Silva',
                                  colaborador=SimpleNamespace(nome='João da Silva'))
        self.assertEqual(nome_usuario(usuario), 'João da Silva')
        usuario.colaborador = SimpleNamespace(nome='  ')
        self.assertEqual(nome_usuario(usuario), 'J Silva')

    def test_sem_grupo_nao_entra(self):
        self.client.force_authenticate(User.objects.create_user('outro', password='x'))
        self.assertEqual(self.client.get('/conformidadeLegal/requisitos/').status_code, 403)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get('/conformidadeLegal/normas/').status_code, 401)

    def test_verificacao_atualiza_situacao_e_proxima_data(self):
        req = self._requisito()
        self.assertEqual(req.situacao, 'NAO_AVALIADO')
        arquivo = SimpleUploadedFile('laudo.pdf', b'%PDF-1.4 teste', content_type='application/pdf')
        r = self.client.post(f'/conformidadeLegal/requisitos/{req.id}/verificar/', {
            'resultado': 'ATENDIDO', 'data': '2026-10-01', 'evidencia': 'Laudo do laboratório.',
            'arquivos': [arquivo],
        }, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(len(r.data['verificacao']['anexos']), 1)
        req.refresh_from_db()
        self.assertEqual((req.situacao, req.ultima_verificacao, req.proxima_verificacao),
                         ('ATENDIDO', date(2026, 10, 1), date(2027, 4, 1)))

        # Apagar a VCL devolve o requisito ao estado anterior.
        vcl = Verificacao.objects.get(requisito=req)
        self.assertEqual(self.client.delete(f'/conformidadeLegal/verificacoes/{vcl.id}/').status_code, 204)
        req.refresh_from_db()
        self.assertEqual((req.situacao, req.proxima_verificacao), ('NAO_AVALIADO', None))

    def test_nao_atendido_exige_plano_e_gera_tarefa_no_kanban(self):
        req = self._requisito()
        url = f'/conformidadeLegal/requisitos/{req.id}/verificar/'
        r = self.client.post(url, {'resultado': 'NAO_ATENDIDO'}, format='multipart')
        self.assertEqual(r.status_code, 400)
        self.assertIn('plano_descricao', r.data)

        quadro = KanbanBoard.objects.create(nome='SGI', criado_por=self.user)
        coluna = KanbanColumn.objects.create(quadro=quadro, titulo='A fazer')
        prazo = (date.today() + timedelta(days=20)).isoformat()
        r = self.client.post(url, {
            'resultado': 'NAO_ATENDIDO', 'plano_descricao': 'Instalar caixa separadora.',
            'plano_prazo': prazo, 'coluna_id': coluna.id,
        }, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        tarefa = KanbanTask.objects.get()
        self.assertEqual((tarefa.coluna_id, tarefa.prioridade, str(tarefa.prazo)), (coluna.id, 'alta', prazo))
        self.assertEqual(r.data['plano']['tarefa'], tarefa.id)
        self.assertEqual(r.data['requisito']['planos_abertos'], 1)

        # Concluir a tarefa no Kanban conclui o plano.
        from django.utils import timezone
        tarefa.concluido_em = timezone.now()
        tarefa.save()
        self.assertEqual(PlanoAcao.objects.get().status, 'CONCLUIDO')
        r = self.client.get('/conformidadeLegal/requisitos/resumo/')
        self.assertEqual(r.data['planos_abertos'], 0)

    def test_quadro_alheio_e_recusado(self):
        req = self._requisito()
        dono = User.objects.create_user('dono', password='x')
        coluna = KanbanColumn.objects.create(quadro=KanbanBoard.objects.create(nome='X', criado_por=dono), titulo='A')
        plano = PlanoAcao.objects.create(requisito=req, descricao='Algo')
        r = self.client.post(f'/conformidadeLegal/planos-acao/{plano.id}/gerar-tarefa/', {'coluna_id': coluna.id},
                             format='json')
        self.assertEqual(r.status_code, 400)
        self.assertFalse(KanbanTask.objects.exists())

    def test_resumo_conta_so_aplicaveis_e_validade_de_licenca(self):
        a = self._requisito()
        self._requisito(aplicabilidade='EM_ANALISE')
        self._requisito(aplicabilidade='NAO_APLICAVEL')
        self.client.post(f'/conformidadeLegal/requisitos/{a.id}/verificar/', {'resultado': 'PARCIAL',
                         'plano_descricao': 'Completar.'}, format='multipart')
        Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='123', ementa='LO da mina',
                             validade=date.today() + timedelta(days=60))
        r = self.client.get('/conformidadeLegal/requisitos/resumo/')
        self.assertEqual((r.data['aplicaveis'], r.data['em_analise'], r.data['nao_aplicaveis']), (1, 1, 1))
        self.assertEqual(r.data['por_situacao']['PARCIAL'], 1)
        self.assertEqual(r.data['indice_conformidade'], 50.0)
        self.assertEqual([v['dias'] for v in r.data['validades']], [60])

    def test_norma_com_requisito_so_sai_com_confirmacao_e_leva_tudo_junto(self):
        import os
        from django.core.files.base import ContentFile
        with tempfile.TemporaryDirectory() as pasta, override_settings(MEDIA_ROOT=pasta):
            req = self._requisito()
            ver = Verificacao.objects.create(requisito=req, data=date(2026, 10, 1), resultado='ATENDIDO')
            anexo = EvidenciaAnexo(verificacao=ver)
            anexo.arquivo.save('foto.jpg', ContentFile(b'x'), save=True)
            caminho = anexo.arquivo.path
            PlanoAcao.objects.create(requisito=req, descricao='Corrigir')
            base = f'/conformidadeLegal/normas/{self.norma.id}/'

            prev = self.client.get(base + 'exclusao/').data
            self.assertEqual((prev['requisitos'], prev['verificacoes'], prev['evidencias'], prev['planos']), (1, 1, 1, 1))

            r = self.client.delete(base)
            self.assertEqual(r.status_code, 409)                  # sem confirmação não apaga
            self.assertEqual(r.data['impacto']['requisitos'], 1)
            self.assertTrue(Norma.objects.filter(pk=self.norma.id).exists())

            self.assertEqual(self.client.delete(base + '?com_requisitos=1').status_code, 204)
            self.assertFalse(Norma.objects.filter(pk=self.norma.id).exists())
            self.assertFalse(Requisito.objects.filter(pk=req.pk).exists())
            self.assertFalse(Verificacao.objects.exists())
            self.assertFalse(PlanoAcao.objects.exists())
            self.assertFalse(EvidenciaAnexo.objects.exists())
            self.assertFalse(os.path.exists(caminho))             # o anexo saiu do disco

    def test_norma_sem_requisito_sai_direto(self):
        self.assertEqual(self.client.delete(f'/conformidadeLegal/normas/{self.norma.id}/').status_code, 204)
        self.assertFalse(Norma.objects.filter(pk=self.norma.id).exists())

    def test_criar_lote(self):
        r = self.client.post('/conformidadeLegal/requisitos/criar-lote/', {'norma': self.norma.id, 'itens': [
            {'descricao': 'A', 'referencia': 'Art. 1', 'origem': 'IA', 'trecho_fonte': 'texto', 'prazo_legal': '2027-03-31'},
            {'descricao': 'B'},
        ]}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(Requisito.objects.filter(norma=self.norma).count(), 2)
        self.assertEqual(str(Requisito.objects.get(descricao='A').proxima_verificacao), '2027-03-31')


class _RespostaFalsa:
    def __init__(self, texto, stop_reason='end_turn'):
        bloco = type('B', (), {'type': 'text', 'text': texto})()
        self.content = [bloco]
        self.stop_reason = stop_reason
        self.usage = type('U', (), {'input_tokens': 1200, 'output_tokens': 300})()


class _ClienteFalso:
    """Imita AnthropicFoundry.messages.stream(...) e guarda os argumentos."""

    def __init__(self, resposta):
        self.resposta = resposta
        self.chamadas = []
        self.messages = self

    def stream(self, **kwargs):
        self.chamadas.append(kwargs)
        resposta = self.resposta

        class _Ctx:
            def __enter__(self_inner):
                return type('S', (), {'get_final_message': lambda _s: resposta})()

            def __exit__(self_inner, *a):
                return False
        return _Ctx()


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class ExtracaoIATests(APITestCase):
    def setUp(self):
        from unittest.mock import patch
        self.patch = patch
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.norma = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='1', ementa='LO',
                                          arquivo=SimpleUploadedFile('lo.pdf', b'%PDF-1.4 conteudo'))
        self.url = f'/conformidadeLegal/normas/{self.norma.id}/extracao-ia/'

    def _rodar(self, cliente, dados=None):
        from django.test import override_settings
        from . import views
        with override_settings(CONFORMIDADE_IA_FOUNDRY_RESOURCE='grupodb', CONFORMIDADE_IA_FOUNDRY_KEY='k'), \
                self.patch('conformidadeLegal.ia._cliente', return_value=cliente), \
                self.patch.object(views.extrair_requisitos, 'delay', side_effect=lambda i: views.extrair_requisitos(i)), \
                self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(self.url, dados or {}, format='json')
        self.assertEqual(r.status_code, 202, r.data)
        return self.client.get(self.url).data

    # Configuração vazia forçada: o .env de quem roda o teste pode ter a chave de verdade.
    @override_settings(CONFORMIDADE_IA_FOUNDRY_RESOURCE='', CONFORMIDADE_IA_FOUNDRY_KEY='')
    def test_sem_configuracao_responde_503(self):
        r = self.client.post(self.url)
        self.assertEqual(r.status_code, 503)
        self.assertFalse(self.client.get(self.url).data['configurada'])

    def test_extracao_grava_sugestoes_normalizadas(self):
        import json
        texto = json.dumps({'resumo': 'LO com 2 condicionantes.', 'itens': [
            {'referencia': 'Condicionante 7', 'descricao': 'Enviar relatório anual', 'tema': 'AMBIENTAL',
             'periodicidade_meses': 12, 'prazo_legal': '2027-03-31', 'trecho': '7. ...até 31 de março', 'pagina': 4,
             'observacao': ''},
            {'referencia': 'x', 'descricao': 'Algo', 'tema': 'INVENTADO', 'periodicidade_meses': 5,
             'prazo_legal': '31/03', 'trecho': 't', 'pagina': 0, 'observacao': ''},
            {'referencia': '', 'descricao': '  ', 'tema': 'SST', 'periodicidade_meses': 0, 'prazo_legal': '',
             'trecho': '', 'pagina': 1, 'observacao': ''},
        ]})
        cliente = _ClienteFalso(_RespostaFalsa(texto))
        dados = self._rodar(cliente)['extracao']
        self.assertEqual(dados['status'], 'CONCLUIDA')
        self.assertEqual(len(dados['itens']), 2)  # o de descrição vazia some
        self.assertEqual(dados['itens'][0]['prazo_legal'], '2027-03-31')
        self.assertEqual((dados['itens'][1]['tema'], dados['itens'][1]['periodicidade_meses'],
                          dados['itens'][1]['prazo_legal'], dados['itens'][1]['pagina']), ('OUTRO', 12, None, None))
        # O PDF vai como documento e a saída é forçada ao esquema.
        chamada = cliente.chamadas[0]
        self.assertEqual(chamada['messages'][0]['content'][0]['type'], 'document')
        self.assertEqual(chamada['output_config']['format']['type'], 'json_schema')

    def test_paginas_recortam_o_pdf_e_voltam_na_numeracao_original(self):
        import base64
        import json
        self.norma.arquivo = SimpleUploadedFile('dt.pdf', _pdf_com_paginas(6))
        self.norma.save()
        # A IA cita a posição no RECORTE: 2ª página do recorte = 4ª do original; 9 não existe.
        texto = json.dumps({'resumo': 'r', 'itens': [
            {'referencia': '4.8', 'descricao': 'Medir', 'tema': 'AMBIENTAL', 'periodicidade_meses': 12,
             'prazo_legal': '', 'trecho': 't', 'pagina': 2, 'observacao': ''},
            {'referencia': '4.9', 'descricao': 'Relatar', 'tema': 'AMBIENTAL', 'periodicidade_meses': 12,
             'prazo_legal': '', 'trecho': 't', 'pagina': 9, 'observacao': ''},
        ]})
        cliente = _ClienteFalso(_RespostaFalsa(texto))
        resposta = self._rodar(cliente, {'paginas': '5, 2;4 - 4'})
        dados = resposta['extracao']
        self.assertEqual(resposta['total_paginas'], 6)
        self.assertEqual(dados['paginas'], '2, 4-5')
        self.assertEqual([i['pagina'] for i in dados['itens']], [4, None])
        # Foi só o recorte, e a instrução avisa que é recorte.
        conteudo = cliente.chamadas[0]['messages'][0]['content']
        enviado = base64.b64decode(conteudo[0]['source']['data'])
        from . import ia
        self.assertEqual(ia.contar_paginas(enviado), 3)
        self.assertIn('RECORTE', conteudo[1]['text'])

    def test_paginas_fora_do_documento_voltam_400(self):
        self.norma.arquivo = SimpleUploadedFile('dt.pdf', _pdf_com_paginas(3))
        self.norma.save()
        with override_settings(CONFORMIDADE_IA_FOUNDRY_RESOURCE='r', CONFORMIDADE_IA_FOUNDRY_KEY='k'):
            r = self.client.post(self.url, {'paginas': '2-7'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('3 página', str(r.data['paginas']))
        self.assertFalse(self.norma.extracoes.exists())

    def test_recusa_vira_erro_legivel(self):
        dados = self._rodar(_ClienteFalso(_RespostaFalsa('', stop_reason='refusal')))['extracao']
        self.assertEqual(dados['status'], 'ERRO')
        self.assertIn('recusou', dados['erro'])

    def test_sem_pdf_e_recusado(self):
        from django.test import override_settings
        self.norma.arquivo = None
        self.norma.save()
        with override_settings(CONFORMIDADE_IA_FOUNDRY_RESOURCE='r', CONFORMIDADE_IA_FOUNDRY_KEY='k'):
            r = self.client.post(self.url)
        self.assertEqual(r.status_code, 400)


def _pdf_com_paginas(n: int) -> bytes:
    """PDF de verdade com n páginas, cada uma escrita "Página k"."""
    import io
    from reportlab.pdfgen import canvas
    saida = io.BytesIO()
    c = canvas.Canvas(saida)
    for k in range(1, n + 1):
        c.drawString(100, 750, f'Pagina {k}')
        c.showPage()
    c.save()
    return saida.getvalue()


class PaginasTests(SimpleTestCase):
    def test_interpreta_intervalos_e_listas(self):
        from .ia import interpretar_paginas
        self.assertEqual(interpretar_paginas('20-21, 24-26', 39), [20, 21, 24, 25, 26])
        self.assertEqual(interpretar_paginas('3; 1 – 2, 3', 5), [1, 2, 3])
        self.assertEqual(interpretar_paginas('  ', 5), [])

    def test_recusa_o_que_nao_da_para_usar(self):
        from .ia import PaginasInvalidas, interpretar_paginas
        for texto in ('a-3', '5-2', '0', '4-6', '2,,x'):
            with self.subTest(texto=texto), self.assertRaises(PaginasInvalidas):
                interpretar_paginas(texto, 5)

    def test_descreve_de_forma_compacta(self):
        from .ia import descrever_paginas
        self.assertEqual(descrever_paginas([1, 2, 3, 7, 9, 10]), '1-3, 7, 9-10')

    def test_recorte_mantem_as_paginas_pedidas_na_ordem(self):
        import io
        from pypdf import PdfReader
        from .ia import recortar_pdf
        recorte = PdfReader(io.BytesIO(recortar_pdf(_pdf_com_paginas(5), [2, 4])))
        self.assertEqual([p.extract_text().strip() for p in recorte.pages], ['Pagina 2', 'Pagina 4'])


class _SessaoDoeFalsa:
    """Imita a API do DOE-RS: busca paginada por dia + detalhe da matéria."""

    def __init__(self, materias, falhar=False):
        self.materias = materias  # {id: dict do detalhe}, com chave extra '_consulta'
        self.falhar = falhar
        self.pedidos = []

    def get(self, url, params=None, headers=None, timeout=None):
        from types import SimpleNamespace
        self.pedidos.append((url, params))
        if self.falhar:
            return SimpleNamespace(status_code=503, json=lambda: {})
        if url.endswith('/materias'):
            chave = params.get('entidade') or params.get('tipoMateria')
            itens = [{'procergs': {'id': i, 'tipo': m['tipoMateria'], 'data': '13/10/2026', 'conteudo': ''}}
                     for i, m in self.materias.items() if m['_consulta'] == chave]
            pagina = itens[(params['page'] - 1) * 10: params['page'] * 10]
            return SimpleNamespace(status_code=200, json=lambda: {'collection': pagina, 'pageSize': 10,
                                                                  'collectionSize': len(itens)})
        i = int(url.rsplit('/', 1)[1])
        return SimpleNamespace(status_code=200, json=lambda: self.materias[i])


def _materia(consulta, tipo, html, orgao='Fundação Estadual de Proteção Ambiental Henrique Luiz Roessler'):
    return {'_consulta': consulta, 'tipoMateria': tipo, 'conteudo': html, 'orgao': orgao,
            'data': '2026-10-13', 'pagina': 99}


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class MonitorDiariosTests(APITestCase):
    FEPAM = 'Fundação Estadual de Proteção Ambiental'
    SEMA = 'Secretaria do Meio Ambiente e Infraestrutura'

    def setUp(self):
        from unittest.mock import patch
        self.patch = patch
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        patch('conformidadeLegal.monitor.PAUSA', 0).start()
        patch('conformidadeLegal.monitor.time.sleep').start()
        self.addCleanup(patch.stopall)

    def _sessao(self):
        return _SessaoDoeFalsa({
            1: _materia(self.FEPAM, 'Atos Administrativos', '<p><b>RESOLUÇÃO Nº 48/2026</b></p><p>Custos do licenciamento corretivo.</p>'),
            2: _materia(self.FEPAM, 'Contratos', '<p>Contrato de limpeza.</p>'),
            3: _materia(self.FEPAM, 'Atos Administrativos', '<p>SÚMULA DE CONTRATO</p><p>Partes: FEPAM e empresa X.</p>'),
            4: _materia(self.SEMA, 'Atos Administrativos', '<p>PORTARIA SEMA Nº 175</p><p>Nomeia servidor.</p>',
                        orgao=self.SEMA),
            5: _materia(self.SEMA, 'Atos Administrativos',
                        '<p>INSTRUÇÃO NORMATIVA SEMA Nº 07/2026</p><p>Áreas de lavra de calcário e emissões atmosféricas.</p>',
                        orgao=self.SEMA),
        })

    def test_coleta_filtra_por_tipo_titulo_e_palavra_e_nao_repete(self):
        from datetime import date
        from . import monitor
        from .models import PublicacaoMonitorada
        sessao = self._sessao()
        execucao = monitor.coletar_doe_rs(date(2026, 10, 13), sessao=sessao)
        self.assertTrue(execucao.ok, execucao.erro)
        # FEPAM entra inteira, menos contrato (tipo) e súmula (título); da SEMA, só o que tem palavra-chave.
        pubs = {p.id_externo: p for p in PublicacaoMonitorada.objects.all()}
        self.assertEqual(set(pubs), {'1', '5'})
        self.assertEqual(pubs['1'].titulo, 'RESOLUÇÃO Nº 48/2026')
        self.assertEqual((pubs['1'].origem, pubs['5'].origem), ('FEPAM', 'SEMA / CONSEMA'))
        self.assertIn('licenciamento', pubs['1'].palavras)
        self.assertEqual(set(pubs['5'].palavras) >= {'calcário', 'lavra', 'emissões atmosféricas'}, True)
        self.assertEqual(pubs['5'].link, 'https://www.diariooficial.rs.gov.br/materia?id=5')
        # Rodar de novo no mesmo dia não duplica nada.
        self.assertEqual(monitor.coletar_doe_rs(date(2026, 10, 13), sessao=self._sessao()).novas, 0)
        self.assertEqual(PublicacaoMonitorada.objects.count(), 2)

    def test_fonte_fora_do_ar_vira_execucao_com_erro_e_aparece_na_situacao(self):
        from datetime import date
        from . import monitor
        execucao = monitor.coletar_doe_rs(date(2026, 10, 13), sessao=_SessaoDoeFalsa({}, falhar=True))
        self.assertFalse(execucao.ok)
        self.assertIn('HTTP 503', execucao.erro)
        r = self.client.get('/conformidadeLegal/novidades/situacao/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual((r.data['fontes'][0]['ok'], r.data['fontes'][0]['falhas_seguidas']), (False, 1))

    def test_dias_a_coletar_recupera_o_que_ficou_para_tras(self):
        from datetime import date
        from . import monitor
        from .models import ExecucaoMonitor
        hoje = date(2026, 10, 13)
        self.assertEqual(monitor.dias_a_coletar('DOE_RS', hoje), [hoje])
        ExecucaoMonitor.objects.create(fonte='DOE_RS', data_referencia=date(2026, 10, 9), ok=True)
        ExecucaoMonitor.objects.create(fonte='DOE_RS', data_referencia=date(2026, 10, 10), ok=False)
        self.assertEqual(monitor.dias_a_coletar('DOE_RS', hoje),
                         [date(2026, 10, 10), date(2026, 10, 11), date(2026, 10, 12), hoje])
        ExecucaoMonitor.objects.create(fonte='DOE_RS', data_referencia=date(2026, 8, 1), ok=True)
        ExecucaoMonitor.objects.filter(data_referencia=date(2026, 10, 9)).delete()
        self.assertEqual(len(monitor.dias_a_coletar('DOE_RS', hoje)), 8)  # no máximo uma semana para trás

    def test_cadastrar_como_norma_preenche_e_anexa_pdf(self):
        from datetime import date
        from .models import PublicacaoMonitorada
        pub = PublicacaoMonitorada.objects.create(
            fonte='DOE_RS', id_externo='1491479', data_publicacao=date(2026, 10, 13),
            orgao='Fundação Estadual de Proteção Ambiental Henrique Luiz Roessler', tipo='Atos Administrativos',
            titulo='PORTARIA FEPAM Nº 634/2026', texto='PORTARIA FEPAM Nº 634/2026\nDispõe sobre o licenciamento.',
            link='https://www.diariooficial.rs.gov.br/materia?id=1491479', pagina=99, tema_sugerido='AMBIENTAL')
        r = self.client.post(f'/conformidadeLegal/novidades/{pub.id}/cadastrar-norma/')
        self.assertEqual(r.status_code, 201, r.data)
        norma = Norma.objects.get(pk=r.data['norma']['id'])
        self.assertEqual((norma.tipo, norma.orgao, norma.numero, norma.ano, norma.esfera),
                         ('PORTARIA', 'FEPAM', '634', 2026, 'ESTADUAL'))
        self.assertEqual(norma.identificacao, 'Portaria FEPAM nº 634/2026')
        with norma.arquivo.open('rb') as f:
            self.assertTrue(f.read(5).startswith(b'%PDF'))
        pub.refresh_from_db()
        self.assertEqual((pub.status, pub.norma_id, pub.avaliado_por_id), ('CADASTRADA', norma.id, self.user.id))
        self.assertEqual(self.client.post(f'/conformidadeLegal/novidades/{pub.id}/cadastrar-norma/').status_code, 400)
        self.assertEqual(self.client.post(f'/conformidadeLegal/novidades/{pub.id}/reabrir/').status_code, 400)

    def test_descartar_e_reabrir(self):
        from datetime import date
        from .models import PublicacaoMonitorada
        pub = PublicacaoMonitorada.objects.create(fonte='DOE_RS', id_externo='9', data_publicacao=date(2026, 10, 13))
        r = self.client.post(f'/conformidadeLegal/novidades/{pub.id}/descartar/', {'motivo': 'Ato interno.'}, format='json')
        self.assertEqual((r.status_code, r.data['status'], r.data['motivo_descarte']), (200, 'DESCARTADA', 'Ato interno.'))
        self.assertEqual(self.client.get('/conformidadeLegal/novidades/?status=NOVA').data, [])
        r = self.client.post(f'/conformidadeLegal/novidades/{pub.id}/reabrir/')
        self.assertEqual((r.data['status'], r.data['avaliado_por_nome']), ('NOVA', None))

    def test_triagem_grava_relevancia_e_motivo(self):
        import json
        from datetime import date
        from . import tasks
        from .models import PublicacaoMonitorada
        pub = PublicacaoMonitorada.objects.create(fonte='DOE_RS', id_externo='5', data_publicacao=date(2026, 10, 13),
                                                  titulo='INSTRUÇÃO NORMATIVA SEMA Nº 07/2026', texto='lavra de calcário')
        resposta = _RespostaFalsa(json.dumps({'relevancia': 'ALTA', 'motivo': 'Muda o licenciamento de lavra.', 'tema': 'MINERARIO'}))
        cliente = type('C', (), {'messages': type('M', (), {'create': lambda _s, **kw: resposta})()})()
        with override_settings(CONFORMIDADE_IA_FOUNDRY_RESOURCE='r', CONFORMIDADE_IA_FOUNDRY_KEY='k'), \
                self.patch('conformidadeLegal.ia._cliente', return_value=cliente):
            tasks.triar_publicacoes()
        pub.refresh_from_db()
        self.assertEqual((pub.relevancia, pub.tema_sugerido, pub.analise), ('ALTA', 'MINERARIO', 'Muda o licenciamento de lavra.'))
        self.assertIsNotNone(pub.analisada_em)


class PalavrasChaveTests(SimpleTestCase):
    def test_casa_sem_acento_e_como_palavra_inteira(self):
        from .monitor import palavras_encontradas
        self.assertEqual(palavras_encontradas('Extração de CALCARIO e emissão atmosférica'), ['calcário', 'emissões atmosféricas'])
        # "cal" solto e "fiscal" não disparam; "brita" sim.
        self.assertEqual(palavras_encontradas('Auditoria fiscal da Cal. Silva'), [])
        self.assertIn('britagem', palavras_encontradas('usina de brita'))


def _xml_dou(id_materia, categoria, tipo, identifica, texto, ementa=''):
    return (f'<xml><article id="{id_materia}" name="x" pubName="DO1" artType="{tipo}" pubDate="14/10/2026" '
            f'artCategory="{categoria}" numberPage="40" '
            f'pdfPage="http://pesquisa.in.gov.br/imprensa/jsp/visualiza/index.jsp?data=14/10/2026&amp;jornal=515&amp;pagina=40" '
            f'idMateria="{id_materia}"><body><Identifica><![CDATA[{identifica}]]></Identifica>'
            f'<Ementa><![CDATA[{ementa}]]></Ementa><Texto><![CDATA[<p>{texto}</p>]]></Texto></body></article></xml>')


class _SessaoInlabsFalsa:
    def __init__(self, zips, login_ok=True, falhas_antes=0):
        # login_ok=False: senha errada (volta para acessar.php). falhas_antes: quantos
        # logins o site "instável" responde sem cookie (HTTP 503) antes de aceitar.
        import requests
        self.cookies = requests.cookies.RequestsCookieJar()
        self.zips, self.login_ok, self.falhas_antes, self.pedidos, self.logins = zips, login_ok, falhas_antes, [], 0

    def post(self, url, data=None, headers=None, timeout=None):
        from types import SimpleNamespace
        self.logins += 1
        if not self.login_ok:
            return SimpleNamespace(status_code=200, url='https://inlabs.in.gov.br/acessar.php')
        if self.logins <= self.falhas_antes:
            return SimpleNamespace(status_code=503, url='https://inlabs.in.gov.br/logar.php')
        self.cookies.set('inlabs_session_cookie', 'abc')
        return SimpleNamespace(status_code=200, url='https://inlabs.in.gov.br/index.php?p=')

    def get(self, url, params=None, headers=None, timeout=None):
        from types import SimpleNamespace
        self.pedidos.append((params, headers))
        if 'dl' not in (params or {}):
            links = ''.join(f'<a href="?p=2026-10-14&amp;dl={n}">baixar</a>' for n in self.zips)
            links += '<a href="?p=2026-10-14&amp;dl=2026-10-14-DO3.zip">s3</a>'  # seção 3 não interessa
            return SimpleNamespace(status_code=200, url='https://inlabs.in.gov.br/index.php?p=2026-10-14', text=links)
        return SimpleNamespace(status_code=200, content=self.zips[params['dl']])


def _zip(*xmls):
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as zf:
        for i, x in enumerate(xmls):
            zf.writestr(f'{i}.xml', '﻿' + x)
        zf.writestr('figura.jpg', b'jpg')
    return buf.getvalue()


@override_settings(MEDIA_ROOT=MEDIA_TESTE, CONFORMIDADE_INLABS_EMAIL='a@b.c', CONFORMIDADE_INLABS_SENHA='x#y')
class MonitorDouTests(APITestCase):
    def setUp(self):
        from unittest.mock import patch
        patch('conformidadeLegal.monitor.time.sleep').start()
        self.addCleanup(patch.stopall)

    def _sessao(self, **kw):
        anm = 'Ministério de Minas e Energia/Agência Nacional de Mineração/Diretoria Colegiada'
        mma = 'Ministério do Meio Ambiente e Mudança do Clima/Secretaria Nacional de Qualidade Ambiental'
        return _SessaoInlabsFalsa({
            '2026-10-14-DO1.zip': _zip(
                _xml_dou('1', anm, 'Resolução', 'RESOLUÇÃO ANM Nº 249, DE 13 DE OUTUBRO DE 2026', 'Regimento da agência.'),
                _xml_dou('2', 'Ministério de Minas e Energia/Agência Nacional de Energia Elétrica', 'Resolução',
                         'RESOLUÇÃO AUTORIZATIVA Nº 16.791', 'Mineração de dados de energia.'),
                _xml_dou('3', mma, 'Portaria', 'PORTARIA SQA/MMA Nº 1.800', 'Nomeia servidor.'),
                _xml_dou('4', mma, 'Portaria', 'PORTARIA SQA/MMA Nº 1.801, DE 7 DE OUTUBRO DE 2026',
                         'Resíduos sólidos da mineração.', ementa='Dispõe sobre resíduos sólidos.'),
                _xml_dou('5', 'Ministério da Previdência Social/Previc', 'Portaria', 'Portaria Previc Nº 756',
                         'Licenciamento de plano de benefícios.'),
                _xml_dou('6', 'Ministério do Trabalho e Emprego/Gabinete do Ministro', 'Portaria', 'PORTARIA MTE Nº 9',
                         'Altera a Norma Regulamentadora nº 22 (NR-22).'),
                _xml_dou('7', anm, 'Despacho', 'DESPACHO', 'Alvará de pesquisa para empresa X.'),
            ),
            '2026-10-14-DO1E.zip': _zip(
                _xml_dou('8', 'Atos do Poder Executivo', 'Decreto', 'DECRETO Nº 12.900', 'Licenciamento ambiental federal.'),
            ),
        }, **kw)

    def test_coleta_dou_filtra_por_orgao_tipo_e_palavra(self):
        from datetime import date
        from . import monitor
        from .models import PublicacaoMonitorada
        sessao = self._sessao()
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=sessao)
        self.assertTrue(e.ok, e.erro)
        self.assertEqual(e.examinadas, 8)
        pubs = {p.id_externo: p for p in PublicacaoMonitorada.objects.filter(fonte='DOU')}
        # ANM entra sem palavra; MMA, MTE e decreto só com palavra; energia, Previc e despacho ficam fora.
        self.assertEqual(set(pubs), {'1', '4', '6', '8'})
        self.assertIn('norma regulamentadora', pubs['6'].palavras)
        self.assertEqual(pubs['4'].texto.split('\n')[0], 'Dispõe sobre resíduos sólidos.')
        self.assertEqual((pubs['4'].pagina, pubs['4'].data_publicacao), (40, date(2026, 10, 14)))
        self.assertIn('pagina=40', pubs['4'].link)
        # A seção 3 não é baixada, e o download leva o cabeçalho "origem" do script oficial.
        baixados = [p['dl'] for p, h in sessao.pedidos if p and 'dl' in p]
        self.assertEqual(sorted(baixados), ['2026-10-14-DO1.zip', '2026-10-14-DO1E.zip'])
        self.assertTrue(all(h.get('origem') for p, h in sessao.pedidos if p and 'dl' in p))
        self.assertEqual(monitor.coletar_dou(date(2026, 10, 14), sessao=self._sessao()).novas, 0)

    def test_login_recusado_vira_falha_legivel(self):
        from datetime import date
        from . import monitor
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=self._sessao(login_ok=False))
        self.assertFalse(e.ok)
        self.assertIn('confira o e-mail e a senha', e.erro)

    def test_login_instavel_tenta_de_novo(self):
        from datetime import date
        from . import monitor
        sessao = self._sessao(falhas_antes=2)
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=sessao)
        self.assertTrue(e.ok, e.erro)
        self.assertEqual(sessao.logins, 3)

    def test_site_fora_nao_culpa_a_senha(self):
        from datetime import date
        from . import monitor
        sessao = self._sessao(falhas_antes=99)
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=sessao)
        self.assertFalse(e.ok)
        self.assertNotIn('senha', e.erro)
        self.assertIn('HTTP 503', e.erro)
        self.assertEqual(sessao.logins, monitor.LOGIN_INLABS_TENTATIVAS)

    @override_settings(CONFORMIDADE_INLABS_SENHA='')
    def test_sem_credencial(self):
        from datetime import date
        from . import monitor
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=self._sessao())
        self.assertIn('sem credencial', e.erro)

    def test_norma_federal_a_partir_do_dou(self):
        from datetime import date
        from . import monitor
        from .models import PublicacaoMonitorada
        pub = PublicacaoMonitorada(fonte='DOU', id_externo='4', data_publicacao=date(2026, 10, 14),
                                   orgao='Ministério do Meio Ambiente e Mudança do Clima/Secretaria Nacional de Qualidade Ambiental',
                                   titulo='PORTARIA SQA/MMA Nº 1.801, DE 7 DE OUTUBRO DE 2026', texto='Resíduos.')
        d = monitor.dados_de_norma(pub)
        self.assertEqual((d['tipo'], d['orgao'], d['numero'], d['ano'], d['esfera']),
                         ('PORTARIA', 'SQA/MMA', '1.801', 2026, 'FEDERAL'))
        self.assertEqual(Norma(**d).identificacao, 'Portaria SQA/MMA nº 1.801/2026')


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class MonitorConfiguravelTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)

    def _semear(self):
        import importlib
        from django.apps import apps
        importlib.import_module('conformidadeLegal.migrations.0004_monitor_configuravel').semear(apps, None)

    def test_carga_inicial_reproduz_o_monitor_de_antes(self):
        from . import monitor
        self._semear()
        textos = ['Lavra de CALCÁRIO e emissões atmosféricas', 'Altera a NR-22 e a Norma Regulamentadora nº 1',
                  'Taxa de licenciamento e outorga de recursos hídricos', 'Auditoria fiscal da Cal. Silva', 'usina de brita']
        do_banco, padrao = monitor.carregar_palavras(), monitor.compilar_palavras(monitor.PALAVRAS_PADRAO)
        self.assertEqual([r for r, _ in do_banco], sorted(r for r, _ in padrao))
        for t in textos:
            self.assertEqual(sorted(monitor.palavras_encontradas(t, do_banco)), sorted(monitor.palavras_encontradas(t, padrao)), t)
        self.assertEqual(monitor.carregar_consultas_doe_rs(), monitor.CONSULTAS_DOE_RS)
        self.assertEqual([(r, t, e, g) for r, t, e, g in monitor.carregar_regras_dou()], monitor.REGRAS_DOU)

    def test_palavra_desativada_e_regex_invalida(self):
        from . import monitor
        self._semear()
        lavra = PalavraChaveMonitor.objects.get(rotulo='lavra')
        r = self.client.patch(f'/conformidadeLegal/monitor/palavras/{lavra.id}/', {'ativo': False}, format='json')
        self.assertEqual(r.status_code, 200)
        self.assertNotIn('lavra', monitor.palavras_encontradas('frente de lavra', monitor.carregar_palavras()))
        r = self.client.post('/conformidadeLegal/monitor/palavras/', {'rotulo': 'x', 'termos': 're:(abc'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('Expressão inválida', str(r.data['termos']))
        r = self.client.post('/conformidadeLegal/monitor/palavras/', {'rotulo': 'pedreira', 'termos': 'Pedreira, pedreiras\n'}, format='json')
        self.assertEqual((r.status_code, r.data['termos']), (201, 'Pedreira\npedreiras'))
        r = self.client.post('/conformidadeLegal/monitor/testar/', {'texto': 'Ampliação da PEDREIRA municipal'}, format='json')
        self.assertIn('pedreira', r.data['palavras'])

    def test_lista_de_palavras_conta_disparos_e_descartes(self):
        from datetime import date
        from .models import PublicacaoMonitorada
        self._semear()
        for i, status in enumerate(['DESCARTADA', 'DESCARTADA', 'CADASTRADA', 'NOVA']):
            PublicacaoMonitorada.objects.create(fonte='DOE_RS', id_externo=str(i), data_publicacao=date.today(),
                                                palavras=['licenciamento'], status=status, origem='FEPAM')
        palavras = {p['rotulo']: p for p in self.client.get('/conformidadeLegal/monitor/palavras/').data}
        self.assertEqual((palavras['licenciamento']['disparos'], palavras['licenciamento']['descartadas'],
                          palavras['licenciamento']['viraram_norma']), (4, 2, 1))
        orgaos = {o['rotulo']: o for o in self.client.get('/conformidadeLegal/monitor/orgaos/').data}
        self.assertEqual(orgaos['FEPAM']['trazidas'], 4)

    def test_orgao_do_dou_so_por_categoria(self):
        r = self.client.post('/conformidadeLegal/monitor/orgaos/', {'fonte': 'DOU', 'rotulo': 'IBAMA', 'campo': 'entidade',
                                                                     'valor': 'IBAMA'}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.post('/conformidadeLegal/monitor/orgaos/', {'fonte': 'DOU', 'rotulo': 'ICMBio', 'campo': 'categoria',
                                                                     'valor': 'Instituto Chico Mendes', 'exige_palavra': True},
                             format='json')
        self.assertEqual(r.status_code, 201, r.data)

    @override_settings(CONFORMIDADE_IA_PERFIL_EMPRESA='do servidor')
    def test_perfil_da_tela_vence_o_do_servidor(self):
        from . import ia
        self.assertEqual(ia.perfil_empresa(), 'do servidor')
        r = self.client.patch('/conformidadeLegal/monitor/configuracao/', {'perfil_empresa': 'Calcário e cal no RS'}, format='json')
        self.assertEqual((r.status_code, r.data['perfil_empresa'], r.data['perfil_servidor']), (200, 'Calcário e cal no RS', 'do servidor'))
        self.assertEqual(ia.perfil_empresa(), 'Calcário e cal no RS')
        self.client.patch('/conformidadeLegal/monitor/configuracao/', {'perfil_empresa': ''}, format='json')
        self.assertEqual(ia.perfil_empresa(), 'do servidor')


class NormasTecnicasTests(SimpleTestCase):
    def test_codigos_padronizados(self):
        from .normas_tecnicas import canonico
        casos = {
            'conforme a ABNT NBR 10004:2004': ('ABNT NBR 10004', '2004'),
            'NBR 10.004/2004': ('ABNT NBR 10004', '2004'),
            'conforme ABNT NBR 10 .004 ou': ('ABNT NBR 10004', ''),  # quebra do pypdf na Diretriz FEPAM
            'pela NBR-9547': ('ABNT NBR 9547', ''),
            'ABNT NBR ISO 14001:2015': ('ABNT NBR ISO 14001', '2015'),
            'NBR ISO/IEC 17025': ('ABNT NBR ISO/IEC 17025', ''),
            'ABNT NBR 15575-1': ('ABNT NBR 15575-1', ''),
        }
        for texto, (codigo, ano) in casos.items():
            item = canonico(texto)
            self.assertEqual((item['codigo'], item['ano']), (codigo, ano), texto)
        self.assertIsNone(canonico('Resolução CONAMA 382'))
        self.assertIsNone(canonico('NBR 10 de algo'))

    def test_termo_do_monitor_casa_as_grafias_e_nao_outro_numero(self):
        from . import monitor
        from .normas_tecnicas import canonico, termos_monitor
        nbr = monitor._padrao([termos_monitor(canonico('ABNT NBR 10004'))])
        for texto in ('ABNT NBR 10004', 'nbr10004', 'NBR 10.004/2004', 'NBR-10004-2', 'Abnt Nbr 10004:2004'):
            self.assertTrue(nbr.search(monitor.normalizar(texto)), texto)
        for texto in ('NBR 100041', 'NBR 1000', 'NBR ISO 10004 é outra'.replace('ISO 10004', 'ISO 9001')):
            self.assertFalse(nbr.search(monitor.normalizar(texto)), texto)
        iso = monitor._padrao([termos_monitor(canonico('NBR ISO 14001'))])
        self.assertTrue(iso.search(monitor.normalizar('certificação ISO 14001:2015')))

    def test_texto_e_ia_se_juntam_por_numero(self):
        from .normas_tecnicas import da_ia, do_texto, juntar
        texto = do_texto([(3, 'Classificar os resíduos pela ABNT NBR 10004:2004 e amostrar\nconforme NBR 10007.'),
                          (5, 'De novo a NBR 10.004.')])
        self.assertEqual([(t['codigo'], t['pagina']) for t in texto], [('ABNT NBR 10004', 3), ('ABNT NBR 10007', 3)])
        ia_ = da_ia([{'codigo': 'NBR 10004', 'titulo': 'Resíduos sólidos — Classificação', 'trecho': '', 'pagina': 0},
                     {'codigo': 'sem código', 'titulo': '', 'trecho': '', 'pagina': 1}])
        juntos = juntar(ia_, texto)
        self.assertEqual([j['codigo'] for j in juntos], ['ABNT NBR 10004', 'ABNT NBR 10007'])
        self.assertEqual((juntos[0]['titulo'], juntos[0]['pagina'], juntos[0]['origem']),
                         ('Resíduos sólidos — Classificação', 3, 'ia+texto'))
        self.assertEqual(juntos[0]['ano'], '2004')


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class NbrNaLeituraTests(ExtracaoIATests):
    def test_leitura_traz_nbrs_e_monitorar_cria_palavra_de_qualquer_orgao(self):
        import json
        from .models import PalavraChaveMonitor
        texto = json.dumps({'resumo': 'LO.', 'itens': [], 'normas_tecnicas': [
            {'codigo': 'ABNT NBR 10004:2004', 'titulo': 'Resíduos sólidos', 'trecho': 'classificar conforme a ABNT NBR 10004', 'pagina': 2},
            {'codigo': 'NBR ISO 14001', 'titulo': '', 'trecho': 'sistema de gestão', 'pagina': 3},
        ]})
        dados = self._rodar(_ClienteFalso(_RespostaFalsa(texto)))['extracao']
        self.assertEqual([(n['codigo'], n['monitorada']) for n in dados['normas_citadas']],
                         [('ABNT NBR 10004', False), ('ABNT NBR ISO 14001', False)])

        url = f'/conformidadeLegal/normas/{self.norma.id}/monitorar-nbr/'
        r = self.client.post(url, {'codigos': ['ABNT NBR 10004', 'ABNT NBR ISO 14001']}, format='json')
        self.assertEqual(r.data['criadas'], ['ABNT NBR 10004', 'ABNT NBR ISO 14001'])
        self.assertTrue(all(n['monitorada'] for n in r.data['extracao']['normas_citadas']))
        p = PalavraChaveMonitor.objects.get(rotulo='ABNT NBR 10004')
        self.assertEqual((p.qualquer_orgao, p.norma_origem_id), (True, self.norma.id))
        # Repetir não duplica; desativada volta a valer.
        p.ativo = False
        p.save()
        r = self.client.post(url, {'codigos': ['NBR 10.004']}, format='json')
        self.assertEqual((r.data['criadas'], r.data['reativadas']), ([], ['ABNT NBR 10004']))
        self.assertEqual(self.client.post(url, {'codigos': ['CONAMA 382']}, format='json').status_code, 400)

    def test_leitura_antiga_ganha_nbrs_pela_regex_sem_ia(self):
        from .models import ExtracaoIA
        self.norma.arquivo = SimpleUploadedFile('lo.pdf', _pdf_com_texto('Condicionante 4: classificar pela ABNT NBR 10004.'))
        self.norma.save()
        ex = ExtracaoIA.objects.create(norma=self.norma, status='CONCLUIDA', itens=[])
        self.assertIsNone(ex.normas_citadas)
        dados = self.client.get(self.url).data['extracao']
        self.assertEqual([(n['codigo'], n['pagina']) for n in dados['normas_citadas']], [('ABNT NBR 10004', 1)])


def _pdf_com_texto(texto: str) -> bytes:
    """PDF de uma página com texto extraível (o pypdf só lê texto de fonte de verdade)."""
    conteudo = f'BT /F1 12 Tf 72 720 Td ({texto}) Tj ET'.encode('latin-1')
    objetos = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length %d >>\nstream\n' % len(conteudo) + conteudo + b'\nendstream',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    saida, offsets = b'%PDF-1.4\n', []
    for i, obj in enumerate(objetos, 1):
        offsets.append(len(saida))
        saida += b'%d 0 obj\n' % i + obj + b'\nendobj\n'
    xref = len(saida)
    saida += b'xref\n0 %d\n0000000000 65535 f \n' % (len(objetos) + 1)
    saida += b''.join(b'%010d 00000 n \n' % o for o in offsets)
    saida += b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF' % (len(objetos) + 1, xref)
    return saida


@override_settings(MEDIA_ROOT=MEDIA_TESTE, CONFORMIDADE_INLABS_EMAIL='a@b.c', CONFORMIDADE_INLABS_SENHA='x#y')
class NbrNoDouTests(APITestCase):
    def setUp(self):
        from unittest.mock import patch
        patch('conformidadeLegal.monitor.time.sleep').start()
        self.addCleanup(patch.stopall)

    def test_nbr_vigiada_entra_de_qualquer_orgao_so_em_ato_normativo(self):
        from datetime import date
        from . import monitor
        from .models import PalavraChaveMonitor, PublicacaoMonitorada
        from .normas_tecnicas import canonico, termos_monitor
        PalavraChaveMonitor.objects.create(rotulo='licenciamento', termos='licenciamento')
        PalavraChaveMonitor.objects.create(rotulo='ABNT NBR 10004', termos=termos_monitor(canonico('NBR 10004')),
                                           qualquer_orgao=True)
        inmetro = 'Ministério do Desenvolvimento, Indústria, Comércio e Serviços/Instituto Nacional de Metrologia'
        sessao = _SessaoInlabsFalsa({'2026-10-14-DO1.zip': _zip(
            _xml_dou('1', inmetro, 'Portaria', 'PORTARIA INMETRO Nº 400', 'Os resíduos seguem a ABNT NBR 10.004:2024.'),
            _xml_dou('2', inmetro, 'Portaria', 'PORTARIA INMETRO Nº 401', 'Licenciamento de laboratório.'),
            _xml_dou('3', inmetro, 'Despacho', 'DESPACHO', 'Cita a NBR 10004.'),
        )})
        e = monitor.coletar_dou(date(2026, 10, 14), sessao=sessao)
        self.assertTrue(e.ok, e.erro)
        pubs = list(PublicacaoMonitorada.objects.filter(fonte='DOU'))
        # Só a portaria que cita a NBR: palavra comum não vale fora dos órgãos acompanhados, e despacho não é norma.
        self.assertEqual([(p.id_externo, p.palavras, p.origem) for p in pubs],
                         [('1', ['ABNT NBR 10004'], monitor.ORIGEM_QUALQUER_ORGAO)])


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class TerceirosTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.url = '/conformidadeLegal/terceiros/'
        self.docs = '/conformidadeLegal/documentos-terceiros/'

    def _terceiro(self, **extra):
        r = self.client.post(self.url, {'razao_social': 'Transportes Pampa Ltda', 'categoria': 'RESIDUOS',
                                        'cnpj': '11222333000181', **extra}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    def _doc(self, terceiro, validade=None, arquivo=True, **extra):
        from datetime import date, timedelta
        dados = {'terceiro': terceiro, 'tipo': 'LICENCA_AMBIENTAL', 'descricao': 'LO transporte Classe I', **extra}
        if validade is not None:
            dados['validade'] = (date.today() + timedelta(days=validade)).isoformat()
        if arquivo:
            dados['arquivo'] = SimpleUploadedFile('lo.pdf', b'%PDF-1.4 lo')
        r = self.client.post(self.docs, dados, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    def test_cnpj_validado_e_formatado(self):
        self.assertEqual(self._terceiro()['cnpj'], '11.222.333/0001-81')
        r = self.client.post(self.url, {'razao_social': 'X', 'cnpj': '11222333000182'}, format='json')
        self.assertIn('cnpj', r.data)

    def test_situacao_documental_pelo_pior_documento_exigido(self):
        t = self._terceiro()['id']
        sit = lambda: self.client.get(f'{self.url}{t}/').data['situacao_documental']
        self.assertEqual(sit(), 'SEM_EXIGENCIA')
        self._doc(t, validade=200)
        self.assertEqual(sit(), 'REGULAR')
        alvara = self._doc(t, validade=10, tipo='ALVARA_FUNCIONAMENTO')
        self.assertEqual(alvara['situacao'], 'A_VENCER')
        self.assertEqual(sit(), 'A_VENCER')
        self._doc(t, arquivo=False, tipo='CTF_IBAMA')            # exigido e ainda não enviado
        self.assertEqual(sit(), 'IRREGULAR')
        self._doc(t, validade=-5, tipo='SEGURO', obrigatorio=False)  # vencido, mas não exigido
        lista = self.client.get(self.url, {'situacao': 'IRREGULAR'}).data
        self.assertEqual([x['id'] for x in lista], [t])

    def test_renovar_guarda_a_versao_anterior(self):
        from datetime import date, timedelta
        t = self._terceiro()['id']
        velho = self._doc(t, validade=-3)
        self.assertEqual(velho['situacao'], 'VENCIDO')
        self.assertEqual(self.client.post(f"{self.docs}{velho['id']}/renovar/", {}, format='multipart').status_code, 400)
        r = self.client.post(f"{self.docs}{velho['id']}/renovar/", {
            'validade': (date.today() + timedelta(days=365)).isoformat(),
            'arquivo': SimpleUploadedFile('lo2.pdf', b'%PDF-1.4 nova')}, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data['situacao'], r.data['descricao']), ('EM_DIA', 'LO transporte Classe I'))
        terceiro = self.client.get(f'{self.url}{t}/').data
        self.assertEqual([d['id'] for d in terceiro['documentos']], [r.data['id']])
        self.assertEqual(terceiro['situacao_documental'], 'REGULAR')
        historico = self.client.get(self.docs, {'terceiro': t, 'historico': 1}).data
        self.assertEqual(len(historico), 2)
        self.assertEqual(self.client.post(f"{self.docs}{velho['id']}/renovar/", {
            'arquivo': SimpleUploadedFile('x.pdf', b'%PDF')}, format='multipart').status_code, 400)

    def test_requisito_e_plano_ligados_ao_terceiro(self):
        t = self._terceiro()['id']
        norma = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', ementa='LO')
        r = self.client.post('/conformidadeLegal/requisitos/', {
            'norma': norma.id, 'descricao': 'Destinar resíduos Classe I só com transportador licenciado',
            'aplicabilidade': 'APLICAVEL', 'terceiros': [t]}, format='json')
        self.assertEqual((r.data['terceiros'], r.data['terceiros_nomes']), ([t], ['Transportes Pampa Ltda']))
        req = r.data['id']
        # Plano aberto na verificação herda o terceiro do requisito.
        v = self.client.post(f'/conformidadeLegal/requisitos/{req}/verificar/', {
            'resultado': 'NAO_ATENDIDO', 'data': '2026-10-14', 'plano_descricao': 'Cobrar a LO renovada'}, format='multipart')
        self.assertEqual(v.status_code, 201, v.data)
        planos = self.client.get('/conformidadeLegal/planos-acao/', {'terceiro': t}).data
        self.assertEqual([p['terceiro_nome'] for p in planos], ['Transportes Pampa Ltda'])
        self.assertEqual(len(self.client.get('/conformidadeLegal/requisitos/', {'terceiro': t}).data), 1)
        terceiro = self.client.get(f'{self.url}{t}/').data
        self.assertEqual((terceiro['total_requisitos'], terceiro['planos_abertos']), (1, 1))
        # Com ligação não se exclui: inativa.
        self.assertEqual(self.client.delete(f'{self.url}{t}/').status_code, 400)

    def test_requisito_com_varios_terceiros(self):
        from .models import Terceiro
        a = self._terceiro()['id']
        b = Terceiro.objects.create(razao_social='Laboratório Sul').id
        norma = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', ementa='LO')
        r = self.client.post('/conformidadeLegal/requisitos/', {
            'norma': norma.id, 'descricao': 'Monitorar efluente', 'aplicabilidade': 'APLICAVEL', 'terceiros': [a, b]}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        req = r.data['id']
        self.assertEqual(sorted(r.data['terceiros_nomes']), ['Laboratório Sul', 'Transportes Pampa Ltda'])
        # Aparece no filtro de cada um, uma vez só.
        for t in (a, b):
            self.assertEqual([x['id'] for x in self.client.get('/conformidadeLegal/requisitos/', {'terceiro': t}).data], [req])
        self.assertEqual(self.client.get('/conformidadeLegal/requisitos/resumo/', {'terceiro': b}).data['aplicaveis'], 1)
        # Com dois terceiros, o plano não adivinha qual: fica sem, até alguém escolher.
        v = self.client.post(f'/conformidadeLegal/requisitos/{req}/verificar/', {
            'resultado': 'NAO_ATENDIDO', 'data': '2026-10-14', 'plano_descricao': 'Trocar o laboratório'}, format='multipart')
        self.assertIsNone(v.data['plano']['terceiro'])
        # Lote troca a lista; lista vazia tira todos.
        self.client.post('/conformidadeLegal/requisitos/editar-lote/', {'ids': [req], 'campos': {'terceiros': [b]}}, format='json')
        self.assertEqual(self.client.get(f'/conformidadeLegal/requisitos/{req}/').data['terceiros'], [b])
        r = self.client.patch(f'/conformidadeLegal/requisitos/{req}/', {'terceiros': []}, format='json')
        self.assertEqual(r.data['terceiros'], [])

    def test_painel_conta_terceiros_e_alertas(self):
        t = self._terceiro()['id']
        self._doc(t, validade=-1)
        self._doc(t, arquivo=False, tipo='ALVARA_FUNCIONAMENTO')
        resumo = self.client.get('/conformidadeLegal/requisitos/resumo/').data['terceiros']
        self.assertEqual((resumo['ativos'], resumo['irregulares'], resumo['documentos_vencidos'], resumo['documentos_pendentes']),
                         (1, 1, 1, 1))
        self.assertEqual([a['situacao'] for a in resumo['alertas']], ['VENCIDO', 'PENDENTE'])

    def test_validade_antes_da_emissao(self):
        t = self._terceiro()['id']
        r = self.client.post(self.docs, {'terceiro': t, 'tipo': 'PPCI', 'emissao': '2026-05-01', 'validade': '2026-04-01'},
                             format='multipart')
        self.assertIn('validade', r.data)


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class AvisosTests(APITestCase):
    def setUp(self):
        from datetime import date, timedelta
        from .models import PlanoAcao, Requisito
        self.hoje = date.today()
        self.user = User.objects.create_user('sgi', password='x', email='sgi@grupodb.com.br', first_name='Ana', last_name='Souza')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.outro = User.objects.create_user('carlos', password='x', email='carlos@grupodb.com.br', first_name='Carlos')
        self.client.force_authenticate(self.user)
        self.norma = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='1', ementa='LO', tema='AMBIENTAL',
                                          validade=self.hoje + timedelta(days=100))
        nr = Norma.objects.create(tipo='NR', orgao='MTE', numero='22', ementa='NR-22', tema='SST')
        R = lambda **kw: Requisito.objects.create(aplicabilidade='APLICAVEL', descricao='x', **kw)
        self.matriz, self.atm = _filiais()
        self.r1 = R(norma=self.norma, tema='AMBIENTAL', responsavel=self.user, unidade=self.matriz,
                    proxima_verificacao=self.hoje - timedelta(days=2))
        self.r2 = R(norma=nr, tema='SST', responsavel=self.outro, unidade=self.atm,
                    proxima_verificacao=self.hoje + timedelta(days=20))
        R(norma=nr, tema='SST', proxima_verificacao=self.hoje + timedelta(days=200))
        PlanoAcao.objects.create(requisito=self.r2, descricao='Treinar', responsavel=self.outro, prazo=self.hoje + timedelta(days=3))

    def test_resumo_filtra_e_usa_os_prazos(self):
        url = '/conformidadeLegal/requisitos/resumo/'
        r = self.client.get(url).data
        self.assertEqual((r['aplicaveis'], r['verificacoes_vencidas'], r['verificacoes_a_vencer']), (3, 1, 1))
        self.assertEqual([u['nome'] for u in r['unidades']], ['Matriz', 'F08 - UP ATM'])
        self.assertEqual(len(r['validades']), 1)  # LO vence em 100 dias, dentro dos 180
        r = self.client.get(url, {'tema': 'SST'}).data
        self.assertEqual((r['aplicaveis'], r['verificacoes_vencidas'], r['planos_abertos'], len(r['validades'])), (2, 0, 1, 0))
        # Matriz: o requisito dela e o que vale para todas (sem unidade); o da ATM fica de fora.
        r = self.client.get(url, {'unidade': self.matriz.id}).data
        self.assertEqual(r['aplicaveis'], 2)
        r = self.client.get(url, {'responsavel': self.outro.id}).data
        self.assertEqual((r['aplicaveis'], r['planos_abertos'], len(r['planos_proximos'])), (1, 1, 1))
        # Prazo de aviso menor: a verificação de 20 dias e a LO de 100 dias saem da lista.
        self.assertEqual(self.client.patch('/conformidadeLegal/avisos/', {'dias_verificacao': 10, 'dias_licenca': 90},
                                           format='json').status_code, 200)
        r = self.client.get(url).data
        self.assertEqual((r['verificacoes_a_vencer'], len(r['validades']), r['prazos']['verificacao']), (0, 0, 10))

    def test_configuracao_valida_emails_e_prazos(self):
        r = self.client.patch('/conformidadeLegal/avisos/', {'email_destinatarios': 'sgi@grupodb.com.br; nao-e-email'}, format='json')
        self.assertIn('email_destinatarios', r.data)
        r = self.client.patch('/conformidadeLegal/avisos/', {'dias_plano': 0}, format='json')
        self.assertIn('dias_plano', r.data)
        r = self.client.patch('/conformidadeLegal/avisos/', {'email_destinatarios': 'a@b.com, c@d.com'}, format='json')
        self.assertEqual(r.data['email_destinatarios'], 'a@b.com\nc@d.com')

    def test_email_completo_para_equipe_e_so_o_seu_para_cada_responsavel(self):
        from django.core import mail
        from . import avisos
        from .models import ConfiguracaoAvisos
        cfg = ConfiguracaoAvisos.atual()
        cfg.email_ativo, cfg.email_destinatarios, cfg.email_frequencia = True, 'meio.ambiente@grupodb.com.br', 'DIARIA'
        cfg.save()
        r = avisos.enviar_avisos(forcar=True)
        self.assertEqual(r['enviados'], 3)
        para = {m.to[0]: m for m in mail.outbox}
        self.assertEqual(set(para), {'meio.ambiente@grupodb.com.br', 'sgi@grupodb.com.br', 'carlos@grupodb.com.br'})
        self.assertIn('LO', para['meio.ambiente@grupodb.com.br'].body)          # licença só no completo
        self.assertNotIn('Licenças', para['carlos@grupodb.com.br'].body)
        self.assertIn('Treinar', para['carlos@grupodb.com.br'].body)
        self.assertNotIn('Treinar', para['sgi@grupodb.com.br'].body)            # o plano é do Carlos
        self.assertIn('2 item(ns)', para['carlos@grupodb.com.br'].subject)
        # Mesmo dia: não manda de novo.
        self.assertTrue(avisos.enviar_avisos().get('pulado'))
        self.assertEqual(len(mail.outbox), 3)

    def test_frequencia_semanal_e_fim_de_semana(self):
        from datetime import datetime
        from django.utils import timezone as tz
        from . import avisos
        from .models import ConfiguracaoAvisos
        cfg = ConfiguracaoAvisos(email_frequencia='SEMANAL', email_dia_semana=2)
        quarta = tz.make_aware(datetime(2026, 10, 14, 8, 0))
        self.assertTrue(avisos.dia_de_enviar(cfg, quarta))
        self.assertFalse(avisos.dia_de_enviar(cfg, tz.make_aware(datetime(2026, 10, 15, 8, 0))))
        cfg.email_frequencia = 'DIARIA'
        self.assertFalse(avisos.dia_de_enviar(cfg, tz.make_aware(datetime(2026, 10, 17, 8, 0))))  # sábado

    def test_desligado_nao_envia_e_teste_manda_para_mim(self):
        from django.core import mail
        from . import avisos
        self.assertTrue(avisos.enviar_avisos().get('pulado'))
        r = self.client.post('/conformidadeLegal/avisos/testar/', {}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual([m.to for m in mail.outbox], [['sgi@grupodb.com.br']])
        prev = self.client.get('/conformidadeLegal/avisos/destinatarios/').data
        self.assertEqual({d['email']: d['itens'] for d in prev['destinatarios']},
                         {'sgi@grupodb.com.br': 1, 'carlos@grupodb.com.br': 2})

    def test_prazo_do_documento_de_terceiro_e_configuravel(self):
        from datetime import timedelta
        from .models import ConfiguracaoAvisos, DocumentoTerceiro, Terceiro, limpar_cache_avisos
        t = Terceiro.objects.create(razao_social='X')
        d = DocumentoTerceiro.objects.create(terceiro=t, tipo='PPCI', validade=self.hoje + timedelta(days=45),
                                             arquivo=SimpleUploadedFile('a.pdf', b'%PDF'))
        self.assertEqual(d.situacao, 'EM_DIA')
        cfg = ConfiguracaoAvisos.atual()
        cfg.dias_documento = 60
        cfg.save()
        limpar_cache_avisos()
        self.assertEqual(d.situacao, 'A_VENCER')


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class AntecedenciaPorItemTests(APITestCase):
    def setUp(self):
        from datetime import date
        self.hoje = date.today()
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)

    def test_item_com_antecedencia_propria_entra_antes(self):
        from datetime import timedelta
        from . import avisos
        from .models import DocumentoTerceiro, Requisito, Terceiro
        lo = Norma.objects.create(tipo='LICENCA', ementa='LO', validade=self.hoje + timedelta(days=200))
        outra = Norma.objects.create(tipo='OUTORGA', ementa='Outorga', validade=self.hoje + timedelta(days=200), dias_aviso=240)
        r_pad = Requisito.objects.create(norma=lo, descricao='a', aplicabilidade='APLICAVEL',
                                         proxima_verificacao=self.hoje + timedelta(days=45))
        r_prop = Requisito.objects.create(norma=lo, descricao='b', aplicabilidade='APLICAVEL', dias_aviso=60,
                                          proxima_verificacao=self.hoje + timedelta(days=45))
        p = avisos.pendencias()
        self.assertEqual([n.id for n in p['licencas']], [outra.id])          # 200 dias: só a de 240 de aviso
        self.assertEqual([r.id for r in p['verificacoes']], [r_prop.id])     # 45 dias: só a de 60 de aviso
        r = self.client.get('/conformidadeLegal/requisitos/resumo/').data
        self.assertEqual((r['verificacoes_a_vencer'], len(r['validades'])), (1, 1))
        # Documento de terceiro: 45 dias, padrão 30 = em dia; com 60 próprios = a vencer.
        t = Terceiro.objects.create(razao_social='X')
        d = DocumentoTerceiro.objects.create(terceiro=t, tipo='PPCI', validade=self.hoje + timedelta(days=45),
                                             arquivo=SimpleUploadedFile('a.pdf', b'%PDF'))
        self.assertEqual(d.situacao, 'EM_DIA')
        d.dias_aviso = 60
        self.assertEqual(d.situacao, 'A_VENCER')
        self.assertFalse(r_pad.id in [x.id for x in p['verificacoes']])

    def test_campo_vai_pela_api(self):
        n = self.client.post('/conformidadeLegal/normas/', {'ementa': 'LO', 'tipo': 'LICENCA', 'dias_aviso': 150}, format='json')
        self.assertEqual(n.data['dias_aviso'], 150)
        r = self.client.post('/conformidadeLegal/requisitos/', {'norma': n.data['id'], 'descricao': 'x', 'dias_aviso': 15}, format='json')
        self.assertEqual(r.data['dias_aviso'], 15)


def _jpeg(largura=40, altura=20, cor=(200, 30, 30)) -> bytes:
    import io
    from PIL import Image
    saida = io.BytesIO()
    Image.new('RGB', (largura, altura), cor).save(saida, format='JPEG')
    return saida.getvalue()


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class RelatorioAnualTests(APITestCase):
    """Relatório Anual de Acompanhamento da licença (prévia, Word, PDF, repetir, legendas)."""

    INICIO, FIM = '2024-12-01', '2025-11-30'

    def setUp(self):
        self.user = User.objects.create_user('meioambiente', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.lo = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='00893', ano=2024, esfera='ESTADUAL',
                                       tema='AMBIENTAL', ementa='LO da lavra de calcário.')

    def _req(self, ref, secao='', **extra):
        return Requisito.objects.create(norma=self.lo, referencia=ref, secao=secao, descricao=f'Cumprir {ref}.',
                                        trecho_fonte=f'Texto literal da condicionante {ref}.', tema='AMBIENTAL',
                                        aplicabilidade=extra.pop('aplicabilidade', 'APLICAVEL'), **extra)

    def _vcl(self, req, data, evidencia, resultado='ATENDIDO', fotos=()):
        vcl = Verificacao.objects.create(requisito=req, data=date.fromisoformat(data), resultado=resultado,
                                         evidencia=evidencia, verificado_por=self.user)
        for nome, legenda in fotos:
            EvidenciaAnexo.objects.create(verificacao=vcl, nome=nome, legenda=legenda,
                                          arquivo=SimpleUploadedFile(nome, _jpeg() if nome.endswith('.jpg') else b'%PDF-1.4'))
        return vcl

    def _previa(self, **params):
        r = self.client.get(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/',
                            {'inicio': self.INICIO, 'fim': self.FIM, **params})
        self.assertEqual(r.status_code, 200, getattr(r, 'data', r))
        return r.data

    def test_ordem_numerica_secoes_e_resposta_do_periodo(self):
        a = self._req('5.10', 'Quanto à Atividade')
        b = self._req('5.2', 'Quanto à Atividade')
        c = self._req('1.1', 'Quanto ao Empreendimento')
        d = self._req('1.2')  # sem seção: fica no capítulo anterior
        self._vcl(b, '2025-03-10', 'Antiga do período.')
        self._vcl(b, '2025-06-10', 'Operação dentro das poligonais.', fotos=[('frente.jpg', 'Frente de lavra'),
                                                                            ('laudo.pdf', '')])
        self._vcl(a, '2024-05-01', 'Resposta do ano passado.')  # fora do período
        self._vcl(c, '2025-01-05', 'Em conformidade.')
        self._req('9.9', aplicabilidade='NAO_APLICAVEL', justificativa='Não há captação de água.')

        dados = self._previa()
        titulos = [s['titulo'] for s in dados['secoes']]
        self.assertEqual(titulos, ['1. Quanto ao Empreendimento', '5. Quanto à Atividade'])
        refs = [i['referencia'] for s in dados['secoes'] for i in s['itens']]
        self.assertEqual(refs, ['1.1', '1.2', '5.2', '5.10', '9.9'])

        itens = {i['referencia']: i for s in dados['secoes'] for i in s['itens']}
        self.assertEqual(itens['5.2']['resposta'], 'Operação dentro das poligonais.')
        self.assertEqual([f['numero'] for f in itens['5.2']['fotos']], [1])
        self.assertEqual([x['nome'] for x in itens['5.2']['documentos']], ['laudo.pdf'])
        self.assertEqual(itens['5.2']['texto'], 'Texto literal da condicionante 5.2.')
        self.assertTrue(itens['5.10']['pendente'])
        self.assertEqual(itens['5.10']['anterior']['evidencia'], 'Resposta do ano passado.')
        self.assertTrue(itens['1.2']['pendente'])
        self.assertIsNone(itens['1.2']['anterior'])
        self.assertEqual((itens['9.9']['origem_resposta'], itens['9.9']['resposta']),
                         ('JUSTIFICATIVA', 'Não há captação de água.'))
        self.assertEqual(dados['totais'], {'itens': 5, 'fora': 0, 'respondidos': 3, 'pendentes': 2, 'com_anterior': 1,
                                           'fotos': 1, 'sem_legenda': 0})
        self.assertIn('dezembro de 2024 e novembro de 2025', dados['apresentacao'])
        del d

    def test_configuracao_salva_e_vazio_volta_ao_padrao(self):
        url = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/?inicio={self.INICIO}&fim={self.FIM}'
        r = self.client.patch(url, {'item': '14', 'licenca': 'LO nº 00893/2024', 'atividade': 'Extração de Calcário',
                                    'desconhecido': 'x'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.lo.refresh_from_db()
        self.assertNotIn('desconhecido', self.lo.relatorio_anual)
        self.assertIn('atendimento do item 14 e seus subitens, das condições e restrições da LO nº 00893/2024',
                      r.data['apresentacao'])
        self.assertEqual(r.data['cabecalho'], 'Relatório Anual de Acompanhamento das Atividades – FEPAM – LO nº 00893/2024')
        r = self.client.patch(url, {'licenca': ''}, format='json')
        self.assertEqual(r.data['config']['licenca'], 'Licença de Operação nº 00893/2024')

    def test_subitem_sem_resposta_propria_nao_fica_pendente(self):
        sub = self._req('3.1.1', 'Quanto à Localização')
        self._req('3.1.5')
        url = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/?inicio={self.INICIO}&fim={self.FIM}'
        r = self.client.patch(url, {'sem_resposta': [sub.id, 'x']}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.lo.refresh_from_db()
        self.assertEqual(self.lo.relatorio_anual['sem_resposta'], [])  # lista inválida não grava lixo
        r = self.client.patch(url, {'sem_resposta': [sub.id]}, format='json')
        itens = {i['referencia']: i for s in r.data['secoes'] for i in s['itens']}
        self.assertEqual((itens['3.1.1']['pendente'], itens['3.1.1']['sem_resposta_propria']), (False, True))
        self.assertTrue(itens['3.1.5']['pendente'])
        self.assertEqual(r.data['totais']['pendentes'], 1)

    def test_word_e_pdf_com_foto_e_pendencia(self):
        import io
        from docx import Document
        req = self._req('1.10', 'Quanto ao Empreendimento')
        self._vcl(req, '2025-04-01', 'Canaletas levam a água às bacias.\nLimpeza bimestral.',
                  fotos=[('canal.jpg', 'Canais de drenagem')])
        self._req('1.11')
        base = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/arquivo/'

        r = self.client.get(base, {'inicio': self.INICIO, 'fim': self.FIM, 'formato': 'docx'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('Relatorio_Anual_Licen', r['Content-Disposition'])
        doc = Document(io.BytesIO(r.content))
        texto = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('1. Quanto ao Empreendimento:', texto)
        self.assertIn('1.10-\tTexto literal da condicionante 1.10.', texto)
        self.assertIn('Limpeza bimestral.', texto)
        self.assertIn('Figura 1: Canais de drenagem.', texto)
        self.assertIn('[Sem resposta registrada no período', texto)
        self.assertEqual(len(doc.inline_shapes), 2)  # logo da capa + a foto

        r = self.client.get(base, {'inicio': self.INICIO, 'fim': self.FIM, 'formato': 'pdf'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))
        from pypdf import PdfReader
        paginas = PdfReader(io.BytesIO(r.content)).pages
        corpo = ' '.join(p.extract_text() for p in paginas)
        self.assertIn('Figura 1: Canais de drenagem.', corpo)
        self.assertIn('RELATÓRIO ANUAL', paginas[0].extract_text())

        r = self.client.get(base, {'formato': 'odt'})
        self.assertEqual(r.status_code, 400)

    def test_foto_sumida_do_disco_nao_derruba_o_relatorio(self):
        import os
        req = self._req('2.1')
        vcl = self._vcl(req, '2025-04-01', 'Ok.', fotos=[('placa.jpg', 'Placa')])
        os.remove(vcl.anexos.first().arquivo.path)
        for formato in ('docx', 'pdf'):
            r = self.client.get(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/arquivo/',
                                {'inicio': self.INICIO, 'fim': self.FIM, 'formato': formato})
            self.assertEqual(r.status_code, 200)

    def test_repetir_resposta_do_ano_anterior(self):
        sim = self._req('1.7')
        com_foto = self._req('1.11')
        nao_atendido = self._req('4.9')
        ja_respondido = self._req('1.1')
        nunca = self._req('1.2')
        self._vcl(sim, '2024-10-01', 'O Empreendedor está ciente dessa obrigatoriedade.')
        self._vcl(com_foto, '2024-10-01', 'Guarita na entrada.', fotos=[('guarita.jpg', 'Guarita')])
        self._vcl(nao_atendido, '2024-10-01', 'Falta plano de monitoramento.', resultado='PARCIAL')
        self._vcl(ja_respondido, '2025-02-01', 'Em conformidade.')

        r = self.client.post(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/repetir/'
                             f'?inicio={self.INICIO}&fim={self.FIM}',
                             {'requisitos': [sim.id, com_foto.id, nao_atendido.id, ja_respondido.id, nunca.id],
                              'data': '2025-11-20', 'fotos': True}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['repetidas'], 2)
        self.assertEqual(sorted(i['referencia'] for i in r.data['ignoradas']), ['1.1', '1.2', '4.9'])
        nova = com_foto.verificacoes.order_by('-data').first()
        self.assertEqual((nova.data, nova.resultado, nova.evidencia), (date(2025, 11, 20), 'ATENDIDO', 'Guarita na entrada.'))
        self.assertIn('repetida da verificação de 01/10/2024', nova.observacao)
        self.assertEqual([a.legenda for a in nova.anexos.all()], ['Guarita'])
        sim.refresh_from_db()
        self.assertEqual(sim.ultima_verificacao, date(2025, 11, 20))
        self.assertEqual(r.data['relatorio']['totais']['pendentes'], 2)

        r = self.client.post(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/repetir/'
                             f'?inicio={self.INICIO}&fim={self.FIM}',
                             {'requisitos': [nunca.id], 'data': '2026-01-10'}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_legendas_so_das_fotos_desta_licenca(self):
        req = self._req('1.11')
        vcl = self._vcl(req, '2025-04-01', 'Placas.', fotos=[('placa.jpg', '')])
        outra = Norma.objects.create(tipo='LICENCA', numero='1', ementa='Outra.')
        alheio = Requisito.objects.create(norma=outra, descricao='x', aplicabilidade='APLICAVEL')
        foto_alheia = EvidenciaAnexo.objects.create(
            verificacao=Verificacao.objects.create(requisito=alheio, resultado='ATENDIDO'),
            nome='y.jpg', arquivo=SimpleUploadedFile('y.jpg', _jpeg()))
        foto = vcl.anexos.first()
        url = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/legendas/'
        r = self.client.post(url, {'legendas': {str(foto.id): 'Placa de advertência', str(foto_alheia.id): 'x'}},
                             format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.post(url, {'legendas': {str(foto.id): 'Placa de advertência'}}, format='json')
        self.assertEqual(r.status_code, 200)
        foto.refresh_from_db()
        self.assertEqual(foto.legenda, 'Placa de advertência')

    def test_secao_entra_no_lote_e_na_leitura_da_ia(self):
        from . import ia
        req = self._req('5.1')
        r = self.client.post('/conformidadeLegal/requisitos/editar-lote/',
                             {'ids': [req.id], 'campos': {'secao': 'Quanto à Atividade'}}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data[0]['secao'], 'Quanto à Atividade')
        item = ia._normalizar({'referencia': '5.1', 'secao': 'Quanto à Atividade:', 'descricao': 'x'})
        self.assertEqual(item['secao'], 'Quanto à Atividade')
        self.assertIn('TODAS as condicionantes numeradas', ia._instrucao(self.lo, []))


class ChaveReferenciaTests(SimpleTestCase):
    def test_ordem(self):
        from .relatorio_anual import chave_referencia
        refs = ['5.10', '5.2', '4.5-1.', '4.5-2.', 'Art. 3', '', '1.16', '5.13.2', '5.13']
        self.assertEqual(sorted(refs, key=chave_referencia),
                         ['1.16', 'Art. 3', '4.5-1.', '4.5-2.', '5.2', '5.10', '5.13', '5.13.2', ''])


def _pdf_relatorio_modelo() -> bytes:
    """
    Relatório no formato do Word da equipe: condicionante em negrito, resposta normal,
    foto com "Figura N: legenda" logo abaixo, cabeçalho e rodapé nas margens.
    """
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    saida = io.BytesIO()
    c = canvas.Canvas(saida, pagesize=A4)
    altura = A4[1]

    def texto(y, t, negrito=False, x=57):
        c.setFont('Helvetica-Bold' if negrito else 'Helvetica', 11)
        c.drawString(x, y, t)

    def margens(n):
        c.setFont('Helvetica', 9)
        c.drawString(530, altura - 40, str(n))
        c.drawString(110, 30, 'Relatório Anual de Acompanhamento das Atividades – FEPAM – LO nº 9/2024')

    texto(700, 'EXTRAÇÃO DE CALCÁRIO', True, 200)
    texto(680, 'PROCESSO ANM 811.029/1970', True, 200)
    texto(400, 'NOVEMBRO 2025', True, 250)
    c.showPage()
    margens(2)
    texto(760, 'APRESENTAÇÃO', True)
    texto(730, 'A empresa comprova o atendimento do item 14 e seus subitens, das condições e restrições da LO nº 9/2024.')
    texto(690, 'ATENDIMENTO DAS CONDIÇÕES E RESTRIÇÕES DA LO Nº 9/2024', True)
    texto(660, '1. Quanto ao Empreendimento:', True)
    texto(630, '1.1- Manter as ARTs atualizadas;', True)
    texto(610, 'ART nº 123, validade até 2028.')
    texto(580, '1.2- A área deverá ser sinalizada e cercada;', True)
    texto(560, 'Há guarita e placas de advertência.')
    c.drawImage(ImageReader(io.BytesIO(_jpeg(300, 200))), 150, 330, width=300, height=200)
    texto(315, 'Figura 1: Guarita na entrada.', x=220)
    texto(285, '1.3- Os limites ficam restritos às poligonais:', True)
    texto(255, '1.3-1. Poligonal útil: 38 ha;', True)
    texto(235, 'O empreendedor respeita as poligonais.')
    margens(2)
    c.showPage()
    c.save()
    return saida.getvalue()


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class ImportarRelatorioTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user('meioambiente', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.lo = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='9', ano=2024, tema='AMBIENTAL', ementa='LO.')
        # Como a IA cadastrou: "Item 1.1", em análise, sem capítulo; e um item que o relatório não traz.
        self.r11 = Requisito.objects.create(norma=self.lo, referencia='Item 1.1', descricao='Manter ARTs.',
                                            trecho_fonte='Manter as ARTs atualizadas;', aplicabilidade='EM_ANALISE')
        self.renov = Requisito.objects.create(norma=self.lo, referencia='Renovação', descricao='Renovar a LO.',
                                              aplicabilidade='EM_ANALISE')
        self.url = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/importar/'

    def _enviar(self, **extra):
        pdf = SimpleUploadedFile('relatorio_2025.pdf', _pdf_relatorio_modelo(), content_type='application/pdf')
        return self.client.post(self.url, {'arquivo': pdf, **extra}, format='multipart')

    def test_leitura_separa_condicionante_resposta_e_figura(self):
        from . import importar_relatorio as ir
        lido = ir.ler(_pdf_relatorio_modelo())
        itens = {i.referencia: i for i in lido['itens']}
        self.assertEqual(list(itens), ['1.1', '1.2', '1.3', '1.3-1'])
        self.assertEqual(itens['1.1'].texto_condicao, 'Manter as ARTs atualizadas;')
        self.assertEqual(itens['1.1'].texto_resposta, 'ART nº 123, validade até 2028.')
        self.assertEqual(itens['1.2'].secao, 'Quanto ao Empreendimento')
        self.assertEqual([(f.numero, f.legenda) for f in itens['1.2'].figuras], [(1, 'Guarita na entrada.')])
        self.assertTrue(itens['1.2'].figuras[0].imagem.startswith(b'\xff\xd8'))  # o JPEG original
        self.assertEqual(itens['1.3'].texto_resposta, '')
        self.assertEqual(itens['1.3-1'].texto_condicao, 'Poligonal útil: 38 ha;')
        self.assertEqual(lido['data'], date(2025, 11, 30))
        self.assertEqual((lido['capa']['item'], lido['capa']['licenca'], lido['capa']['atividade']),
                         ('14', 'LO nº 9/2024', 'Extração de Calcário'))

    def test_previa_nao_grava(self):
        r = self._enviar()
        self.assertEqual(r.status_code, 200, r.data)
        t = r.data['totais']
        self.assertEqual((t['condicionantes'], t['casadas'], t['novas'], t['figuras'], t['ausentes'], t['viram_aplicavel']),
                         (4, 1, 3, 1, 1, 1))
        self.assertEqual(Verificacao.objects.count(), 0)
        self.assertEqual(Requisito.objects.filter(norma=self.lo).count(), 2)

    def test_aplicar_grava_respostas_fotos_e_e_idempotente(self):
        r = self._enviar(aplicar='1')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data['respondidos'], r.data['fotos'], r.data['criados'], r.data['fora']), (3, 1, 3, 1))
        self.assertEqual(r.data['data'], '2025-11-30')
        self.r11.refresh_from_db()
        self.assertEqual((self.r11.aplicabilidade, self.r11.secao, self.r11.situacao),
                         ('APLICAVEL', 'Quanto ao Empreendimento', 'ATENDIDO'))
        vcl = self.r11.verificacoes.get()
        self.assertEqual((vcl.data, vcl.evidencia), (date(2025, 11, 30), 'ART nº 123, validade até 2028.'))
        r12 = Requisito.objects.get(norma=self.lo, referencia='1.2')
        self.assertEqual((r12.origem, r12.trecho_fonte), ('RELATORIO', 'A área deverá ser sinalizada e cercada;'))
        foto = EvidenciaAnexo.objects.get(verificacao__requisito=r12)
        self.assertEqual(foto.legenda, 'Guarita na entrada.')
        self.lo.refresh_from_db()
        r13 = Requisito.objects.get(norma=self.lo, referencia='1.3')
        self.assertEqual(self.lo.relatorio_anual['sem_resposta'], [r13.id])  # respondida no 1.3-1
        self.assertEqual(self.lo.relatorio_anual['fora'], [self.renov.id])
        self.assertEqual((self.lo.relatorio_anual['item'], self.lo.relatorio_anual['licenca']), ('14', 'LO nº 9/2024'))
        # O relatório do período da importação sai completo.
        rel = self.client.get(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/',
                              {'inicio': '2024-12-01', 'fim': '2025-11-30'}).data
        self.assertEqual((rel['totais']['pendentes'], rel['totais']['fora'], rel['totais']['fotos']), (0, 1, 1))
        # De novo, mesmo PDF e mesma data: nada duplica.
        r = self._enviar(aplicar='1')
        self.assertEqual((r.data['respondidos'], r.data['repetidos'], r.data['criados']), (0, 3, 0))
        self.assertEqual(Verificacao.objects.count(), 3)

    def test_pdf_sem_texto_e_recusado(self):
        r = self.client.post(self.url, {'arquivo': SimpleUploadedFile('x.pdf', b'nao e pdf')}, format='multipart')
        self.assertEqual(r.status_code, 400)

    def test_comando_de_gestao(self):
        import os
        from io import StringIO
        from django.core.management import call_command
        caminho = os.path.join(MEDIA_TESTE, 'rel.pdf')
        with open(caminho, 'wb') as f:
            f.write(_pdf_relatorio_modelo())
        saida = StringIO()
        call_command('importar_relatorio_anual', norma=self.lo.id, pdf=caminho, stdout=saida)
        self.assertIn('Prévia apenas', saida.getvalue())
        self.assertEqual(Verificacao.objects.count(), 0)
        call_command('importar_relatorio_anual', norma=self.lo.id, pdf=caminho, aplicar=True, stdout=saida)
        self.assertIn('3 respostas, 1 fotos, 3 condicionantes criadas', saida.getvalue())


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class UnidadeTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.matriz, self.atm = _filiais()
        self.lo_matriz = Norma.objects.create(tipo='LICENCA', numero='893', ementa='LO matriz', unidade=self.matriz)
        self.lo_atm = Norma.objects.create(tipo='LICENCA', numero='77', ementa='LO ATM', unidade=self.atm)
        self.lei = Norma.objects.create(tipo='LEI', numero='12305', ementa='PNRS')

    def test_lista_de_unidades_com_matriz_primeiro(self):
        r = self.client.get('/conformidadeLegal/unidades/')
        self.assertEqual([u['nome'] for u in r.data], ['Matriz', 'F08 - UP ATM'])

    def test_normas_da_unidade_mais_as_gerais(self):
        r = self.client.get('/conformidadeLegal/normas/', {'unidade': self.atm.id})
        self.assertEqual(sorted(n['ementa'] for n in r.data), ['LO ATM', 'PNRS'])
        self.assertEqual(len(self.client.get('/conformidadeLegal/normas/').data), 3)

    def test_requisito_herda_a_unidade_da_licenca(self):
        r = self.client.post('/conformidadeLegal/requisitos/', {
            'norma': self.lo_atm.id, 'descricao': 'Monitorar.', 'unidade': self.matriz.id}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data['unidade'], r.data['unidade_nome']), (self.atm.id, 'F08 - UP ATM'))
        # Lei geral: a obrigação pode ser só de uma unidade.
        r = self.client.post('/conformidadeLegal/requisitos/', {
            'norma': self.lei.id, 'descricao': 'PGRS.', 'unidade': self.matriz.id}, format='json')
        self.assertEqual(r.data['unidade'], self.matriz.id)
        ids = {x['descricao'] for x in self.client.get('/conformidadeLegal/requisitos/', {'unidade': self.atm.id}).data}
        self.assertEqual(ids, {'Monitorar.'})

    def test_trocar_a_unidade_da_licenca_leva_as_condicionantes(self):
        req = Requisito.objects.create(norma=self.lei, descricao='x')
        self.assertIsNone(req.unidade_id)
        r = self.client.patch(f'/conformidadeLegal/normas/{self.lei.id}/', {'unidade': self.atm.id}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        req.refresh_from_db()
        self.assertEqual(req.unidade_id, self.atm.id)

    def test_terceiros_e_planos_por_unidade(self):
        from .models import Terceiro
        Terceiro.objects.create(razao_social='Coleta Matriz', unidade=self.matriz)
        Terceiro.objects.create(razao_social='Coleta Geral')
        nomes = [t['razao_social'] for t in self.client.get('/conformidadeLegal/terceiros/', {'unidade': self.atm.id}).data]
        self.assertEqual(nomes, ['Coleta Geral'])
        req = Requisito.objects.create(norma=self.lo_matriz, descricao='x', aplicabilidade='APLICAVEL')
        PlanoAcao.objects.create(requisito=req, descricao='Plano matriz')
        self.assertEqual(len(self.client.get('/conformidadeLegal/planos-acao/', {'unidade': self.atm.id}).data), 0)
        self.assertEqual(len(self.client.get('/conformidadeLegal/planos-acao/', {'unidade': self.matriz.id}).data), 1)

    def test_municipio_da_capa_vem_da_unidade(self):
        r = self.client.get(f'/conformidadeLegal/normas/{self.lo_atm.id}/relatorio-anual/')
        self.assertEqual(r.data['config']['municipio'], 'Almirante Tamandaré – PR')


@override_settings(MEDIA_ROOT=MEDIA_TESTE)
class RelatorioForaELogoTests(RelatorioAnualTests):
    """Checkbox "vai para o relatório", número sem o "Item" e logo da capa."""

    def test_fora_do_relatorio_nao_sai_no_documento(self):
        import io
        from docx import Document
        dentro = self._req('Item 1.1', 'Quanto ao Empreendimento')
        fora = self._req('Disposições finais')
        self._vcl(dentro, '2025-04-01', 'Em conformidade.')
        url = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/?inicio={self.INICIO}&fim={self.FIM}'
        r = self.client.patch(url, {'fora': [fora.id]}, format='json')
        itens = {i['requisito']: i for s in r.data['secoes'] for i in s['itens']}
        self.assertEqual((itens[fora.id]['no_relatorio'], itens[fora.id]['pendente']), (False, False))
        self.assertEqual(itens[dentro.id]['referencia'], '1.1')
        self.assertEqual((r.data['totais']['itens'], r.data['totais']['fora'], r.data['totais']['pendentes']), (1, 1, 0))
        arq = self.client.get(f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/arquivo/',
                              {'inicio': self.INICIO, 'fim': self.FIM, 'formato': 'docx'})
        texto = '\n'.join(p.text for p in Document(io.BytesIO(arq.content)).paragraphs)
        self.assertIn('1.1-\tTexto literal da condicionante Item 1.1.', texto)
        self.assertNotIn('Disposições finais', texto)

    def test_logo_proprio_padrao_e_sem_logo(self):
        import io
        from docx import Document
        base = f'/conformidadeLegal/normas/{self.lo.id}/relatorio-anual/'
        arquivo = lambda: self.client.get(base + 'arquivo/', {'formato': 'docx'})  # noqa: E731
        self.assertEqual(len(Document(io.BytesIO(arquivo().content)).inline_shapes), 1)  # logo do grupo
        r = self.client.post(base + 'logo/', {'arquivo': SimpleUploadedFile('logo.png', _jpeg(), content_type='image/png')},
                             format='multipart')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(r.data['logo_proprio'])
        self.assertEqual(len(Document(io.BytesIO(arquivo().content)).inline_shapes), 1)
        self.assertTrue(self.client.get(base + 'arquivo/', {'formato': 'pdf'}).content.startswith(b'%PDF'))
        r = self.client.delete(base + 'logo/')
        self.assertIsNone(r.data['logo_proprio'])
        self.client.patch(base, {'logo': 'nao'}, format='json')
        self.assertEqual(len(Document(io.BytesIO(arquivo().content)).inline_shapes), 0)


class MatrizTests(APITestCase):
    """Matriz de requisitos (planilha do SGI) — importação, vínculo automático e situação calculada."""
    CAB = ['Atividade', 'Obrigações ambientais aplicáveis', 'Origem / referência a confirmar', 'Exemplos de atendimento / controles',
           'Evidências de atendimento', 'Responsável sugerido', 'Frequência / prazo', 'Situação', 'Última verificação', 'Pendências / ações']

    def setUp(self):
        self.user = User.objects.create_user('sgi', password='x')
        self.user.groups.add(Group.objects.create(name='ConformidadeLegal'))
        self.client.force_authenticate(self.user)
        self.pnrs = Norma.objects.create(tipo='LEI', numero='12.305', ano=2010, ementa='Política Nacional de Resíduos Sólidos.')
        self.conama = Norma.objects.create(tipo='RESOLUCAO', orgao='CONAMA', numero='430', ano=2011, ementa='Lançamento de efluentes.')
        self.lo = Norma.objects.create(tipo='LICENCA', orgao='FEPAM', numero='LO 100', ano=2024, ementa='Licença de operação da lavra.',
                                       validade=date.today() + timedelta(days=400))
        self.r_mtr = Requisito.objects.create(norma=self.pnrs, descricao='Emitir MTR para toda saída de resíduo.', aplicabilidade='APLICAVEL')
        self.r_efl = Requisito.objects.create(norma=self.conama, descricao='Monitorar o efluente da caixa separadora.', aplicabilidade='APLICAVEL')
        self.r_lavra = Requisito.objects.create(norma=self.lo, descricao='Manter a frente de lavra sinalizada.', aplicabilidade='APLICAVEL')
        self.r_fora = Requisito.objects.create(norma=self.lo, descricao='Instalar hidrômetro.', aplicabilidade='NAO_APLICAVEL')

    def _planilha(self, linhas):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Matriz de Requisitos'
        ws.append(self.CAB)
        for l in linhas:
            ws.append(l + ['A preencher'] * (len(self.CAB) - len(l)))
        buf = io.BytesIO()
        wb.save(buf)
        return SimpleUploadedFile('matriz.xlsx', buf.getvalue(),
                                  content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    LINHAS = [
        ['Gestão de resíduos sólidos', 'Classificação, destinação adequada; MTR quando exigido', 'Lei 12.305/2010; PGRS', 'Coletores', 'MTR', 'Meio Ambiente', 'Conforme geração'],
        ['Tratamento e lançamento de efluentes', 'Atender padrões de lançamento', 'Resolução CONAMA 430/2011', 'Coletas', 'Laudos', 'Manutenção', 'Conforme licença'],
        ['Mineração a céu aberto', 'Licenciamento ambiental; atendimento às condicionantes', 'Licença ambiental vigente', 'Inspecionar', 'Licença', 'Mineração', 'Conforme a licença'],
        ['Silvicultura, plantio e corte', 'Verificar necessidade de autorização florestal', 'Legislação florestal', 'Mapear', 'Mapas', 'Meio Ambiente', 'Antes'],
    ]

    def _importar(self, linhas=None):
        r = self.client.post('/conformidadeLegal/matriz/importar/', {'arquivo': self._planilha(linhas or self.LINHAS)}, format='multipart')
        self.assertEqual(r.status_code, 201, r.data)
        return r.data

    def _item(self, atividade):
        from .models import MatrizItem
        mi = MatrizItem.objects.get(atividade=atividade)
        return self.client.get(f'/conformidadeLegal/matriz/{mi.id}/').data

    def test_importa_liga_ao_que_existe_e_calcula_a_situacao(self):
        d = self._importar()
        self.assertEqual((d['criados'], d['atualizados']), (4, 0))
        res = self._item('Gestão de resíduos sólidos')
        self.assertIn(self.r_mtr.id, [v['requisito'] for v in res['vinculos']])
        self.assertEqual(res['situacao'], 'NAO_VERIFICADO')           # ligado, mas ninguém verificou ainda
        efl = self._item('Tratamento e lançamento de efluentes')
        self.assertEqual([v['requisito'] for v in efl['vinculos']], [self.r_efl.id])
        lavra = self._item('Mineração a céu aberto')
        self.assertIn(self.lo.id, [v['norma'] for v in lavra['vinculos_normas']])
        self.assertIn(self.r_lavra.id, [v['requisito'] for v in lavra['vinculos']])
        self.assertNotIn(self.r_fora.id, [v['requisito'] for v in lavra['vinculos']])   # não aplicável fica de fora
        self.assertEqual(self._item('Silvicultura, plantio e corte')['situacao'], 'NAO_MAPEADO')

        self.r_mtr.verificacoes.create(resultado='ATENDIDO', data=date.today())
        self.r_mtr.recalcular_situacao()
        res = self._item('Gestão de resíduos sólidos')
        self.assertEqual((res['situacao'], res['ultima_verificacao']), ('CONFORME', date.today().isoformat()))
        self.r_efl.verificacoes.create(resultado='NAO_ATENDIDO', data=date.today())
        self.r_efl.recalcular_situacao()
        efl = self._item('Tratamento e lançamento de efluentes')
        self.assertEqual(efl['situacao'], 'NAO_CONFORME')
        self.assertTrue(any(p['tipo'] == 'NAO_ATENDIDO' for p in efl['pendencias']))
        resumo = self.client.get('/conformidadeLegal/matriz/resumo/').data
        self.assertEqual(resumo['contagem']['CONFORME'], 1)
        self.assertEqual(resumo['total'], 4)

    def test_licenca_vencida_deixa_nao_conforme(self):
        self._importar()
        self.lo.validade = date.today() - timedelta(days=1)
        self.lo.save()
        self.r_lavra.verificacoes.create(resultado='ATENDIDO', data=date.today())
        self.r_lavra.recalcular_situacao()
        lavra = self._item('Mineração a céu aberto')
        self.assertEqual(lavra['situacao'], 'NAO_CONFORME')
        self.assertTrue(any(p['tipo'] == 'LICENCA_VENCIDA' for p in lavra['pendencias']))

    def test_reimportar_atualiza_sem_duplicar_e_vinculo_tirado_nao_volta(self):
        self._importar()
        res = self._item('Gestão de resíduos sólidos')
        r = self.client.delete(f"/conformidadeLegal/matriz/{res['id']}/vinculos/?requisito={self.r_mtr.id}")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn(self.r_mtr.id, [v['requisito'] for v in r.data['vinculos']])
        linhas = [l[:] for l in self.LINHAS]
        linhas[0][1] = 'Texto novo da obrigação de resíduos'
        d = self._importar(linhas)
        self.assertEqual((d['criados'], d['atualizados']), (0, 4))
        res = self._item('Gestão de resíduos sólidos')
        self.assertEqual(res['obrigacoes'], 'Texto novo da obrigação de resíduos')
        self.assertNotIn(self.r_mtr.id, [v['requisito'] for v in res['vinculos']])
        r = self.client.post(f"/conformidadeLegal/matriz/{res['id']}/vinculos/", {'requisito': self.r_mtr.id}, format='json')
        self.assertIn(self.r_mtr.id, [v['requisito'] for v in r.data['vinculos']])

    def test_situacao_manual_exige_justificativa_e_exporta_planilha(self):
        self._importar()
        silv = self._item('Silvicultura, plantio e corte')
        url = f"/conformidadeLegal/matriz/{silv['id']}/"
        self.assertEqual(self.client.patch(url, {'situacao_manual': 'NAO_APLICAVEL'}, format='json').status_code, 400)
        r = self.client.patch(url, {'situacao_manual': 'NAO_APLICAVEL', 'justificativa': 'Sem silvicultura na unidade.'}, format='json')
        self.assertEqual((r.status_code, r.data['situacao'], r.data['situacao_automatica']), (200, 'NAO_APLICAVEL', 'NAO_MAPEADO'))
        from openpyxl import load_workbook
        r = self.client.get('/conformidadeLegal/matriz/exportar/')
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        linhas = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(linhas['Silvicultura, plantio e corte'][7], 'Não aplicável (manual)')

    def test_importar_recusa_planilha_sem_cabecalho(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(['Nada', 'aqui'])
        buf = io.BytesIO()
        wb.save(buf)
        r = self.client.post('/conformidadeLegal/matriz/importar/', {'arquivo': SimpleUploadedFile('x.xlsx', buf.getvalue())}, format='multipart')
        self.assertEqual(r.status_code, 400)
