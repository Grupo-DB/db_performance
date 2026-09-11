"""
Grava a classificação das análises que voltaram preenchidas na planilha.

Em 27/08/2026 exportamos as 585 análises (de 975) que estavam sem
`classificacao`, com a coluna a preencher à mão. Em 28/08 a planilha voltou com
573 respondidas — 12 ficaram em branco, todas de Argamassa/Aditivo, onde a escala
não descreve bem o ensaio.

As respostas moram no JSON ao lado (`management/dados/`) em vez de exigir que a
planilha seja copiada para o servidor: assim a carga viaja junto com o `git pull`,
fica revisável no diff e pode ser repetida sem depender de um arquivo no
`Downloads` de alguém. Para uma rodada nova, use `--planilha`.

Segunda rodada, 09/09/2026: outras 89 análises apareceram sem classificação (a
base cresceu desde agosto). A carga é `classificacao_2026_09_08.json` — 88 com
valor e 1 sem, a análise 184, cuja OS também está sem classificação.

    python manage.py importa_classificacao_analises --dados classificacao_2026_09_08.json
    python manage.py importa_classificacao_analises --dados classificacao_2026_09_08.json --aplicar

Não grava nada sem `--aplicar`.
"""
import json
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

from controleQualidade.analise.models import Analise

# A escala combinada com o laboratório. `Analise.classificacao` é CharField livre,
# sem `choices` — nada no banco impede um valor torto, então a validação é aqui.
CLASSIFICACOES_VALIDAS = ('Simples', 'Parcial', 'Complexa', 'Super Complexa')

DIRETORIO_DADOS = Path(__file__).resolve().parent.parent / 'dados'
DADOS_PADRAO = DIRETORIO_DADOS / 'classificacao_2026_08_28.json'


class Command(BaseCommand):
    help = 'Aplica nas análises a classificação preenchida na planilha de 28/08/2026.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--aplicar', action='store_true',
            help='Grava as classificações. Sem isto, só relata.',
        )
        parser.add_argument(
            '--dados', default='',
            help='Outro JSON de management/dados/ (rodada nova). Aceita o nome do '
                 'arquivo ou o caminho completo. Ex.: '
                 '--dados classificacao_2026_09_08.json (88 análises da base de 08/09).',
        )
        parser.add_argument(
            '--planilha', default='',
            help='Caminho de um .xlsx no lugar do JSON embutido (rodada nova). '
                 'Espera as colunas da exportação: nº da amostra em A, '
                 'classificação em B e ID da análise em Q.',
        )
        parser.add_argument(
            '--sobrescrever', action='store_true',
            help='Também troca a classificação de análise que JÁ tem uma. '
                 'Por padrão elas são preservadas.',
        )

    def handle(self, *args, **opcoes):
        aplicar = opcoes['aplicar']
        itens = (self._le_planilha(opcoes['planilha']) if opcoes['planilha']
                 else self._le_json(opcoes['dados']))

        invalidos = [r for r in itens if r['classificacao'] not in CLASSIFICACOES_VALIDAS]
        if invalidos:
            raise CommandError(
                'Classificação fora da escala em {} linha(s): {}'.format(
                    len(invalidos),
                    ', '.join(f"análise {r['analise_id']}={r['classificacao']!r}"
                              for r in invalidos[:10]),
                )
            )

        if not aplicar:
            self.stdout.write(self.style.WARNING(
                'Simulação (use --aplicar para gravar).\n'))

        atuais = dict(Analise.objects
                      .filter(id__in=[r['analise_id'] for r in itens])
                      .values_list('id', 'classificacao'))

        gravadas, preservadas, inexistentes, iguais = 0, [], [], 0
        for registro in itens:
            analise_id = registro['analise_id']
            nova = registro['classificacao']
            if analise_id not in atuais:
                inexistentes.append(registro)
                continue
            atual = (atuais[analise_id] or '').strip()
            if atual == nova:
                iguais += 1
                continue
            if atual and not opcoes['sobrescrever']:
                # Alguém pode ter classificado pela tela desde a exportação, e o
                # que a pessoa fez olhando o ensaio vale mais que a planilha.
                preservadas.append((registro, atual))
                continue
            if aplicar:
                # `update()` de queryset, e NÃO `save()`: `Analise.data` é
                # `auto_now=True`, ou seja "última alteração". Um `save()` aqui
                # datava 573 análises de hoje e reescrevia esse histórico em
                # massa por causa de um preenchimento administrativo.
                Analise.objects.filter(pk=analise_id).update(classificacao=nova)
            gravadas += 1

        self._relata(itens, gravadas, iguais, preservadas, inexistentes, aplicar)

    # ── Fontes ───────────────────────────────────────────────────────────────

    def _le_json(self, nome=''):
        """Carga embutida. Sem `--dados`, a de 28/08; com, a que for pedida.

        Nome solto (sem barra) é procurado em management/dados/ — é onde as cargas
        moram, e digitar o caminho inteiro na VM só rende erro de digitação.
        """
        caminho = DADOS_PADRAO
        if nome:
            caminho = Path(nome).expanduser()
            if not caminho.is_absolute() and len(caminho.parts) == 1:
                caminho = DIRETORIO_DADOS / nome
        if not caminho.exists():
            raise CommandError(f'Arquivo de dados não encontrado: {caminho}')
        dados = json.loads(caminho.read_text(encoding='utf-8'))
        self.stdout.write(
            f"Fonte: {dados.get('origem')} "
            f"(preenchida em {dados.get('preenchido_em')}) — "
            f"{len(dados['itens'])} respondidas, "
            f"{len(dados.get('sem_resposta') or [])} em branco.\n"
        )
        return dados['itens']

    def _le_planilha(self, caminho):
        try:
            import openpyxl
        except ImportError as erro:                       # pragma: no cover
            raise CommandError('openpyxl não está instalado neste ambiente.') from erro

        planilha = Path(caminho).expanduser()
        if not planilha.exists():
            raise CommandError(f'Planilha não encontrada: {planilha}')

        wb = openpyxl.load_workbook(planilha, read_only=True, data_only=True)
        aba = wb['Sem classificação'] if 'Sem classificação' in wb.sheetnames else wb.worksheets[0]
        itens = []
        for linha in aba.iter_rows(values_only=True):
            analise_id, valor = linha[16] if len(linha) > 16 else None, linha[1]
            # Pula cabeçalho e linha de instrução: só interessa linha cujo ID é
            # número de verdade.
            if not isinstance(analise_id, int):
                continue
            if isinstance(valor, str) and valor.strip():
                itens.append({'analise_id': analise_id, 'amostra': linha[0],
                              'classificacao': valor.strip()})
        wb.close()
        self.stdout.write(f'Fonte: {planilha.name} — {len(itens)} respondidas.\n')
        return itens

    # ── Relatório ────────────────────────────────────────────────────────────

    def _relata(self, itens, gravadas, iguais, preservadas, inexistentes, aplicar):
        verbo = 'Gravadas' if aplicar else 'A gravar'
        self.stdout.write(self.style.SUCCESS(f'{verbo}: {gravadas}'))
        if iguais:
            self.stdout.write(f'Já estavam com a mesma classificação: {iguais}')

        contagem = {}
        for registro in itens:
            contagem[registro['classificacao']] = contagem.get(registro['classificacao'], 0) + 1
        for nome in CLASSIFICACOES_VALIDAS:
            if contagem.get(nome):
                self.stdout.write(f'  {nome}: {contagem[nome]}')

        if preservadas:
            self.stdout.write(self.style.WARNING(
                f'\n{len(preservadas)} análise(s) já classificada(s) no sistema — '
                f'mantidas como estão (use --sobrescrever para trocar):'))
            for registro, atual in preservadas[:20]:
                self.stdout.write(
                    f"  análise {registro['analise_id']} ({registro.get('amostra')}): "
                    f"sistema={atual!r}, planilha={registro['classificacao']!r}")
            if len(preservadas) > 20:
                self.stdout.write(f'  ... e outras {len(preservadas) - 20}')

        if inexistentes:
            self.stdout.write(self.style.ERROR(
                f'\n{len(inexistentes)} análise(s) da planilha não existe(m) mais '
                f'(apagadas depois da exportação):'))
            for registro in inexistentes[:20]:
                self.stdout.write(
                    f"  análise {registro['analise_id']} ({registro.get('amostra')})")
            if len(inexistentes) > 20:
                self.stdout.write(f'  ... e outras {len(inexistentes) - 20}')
